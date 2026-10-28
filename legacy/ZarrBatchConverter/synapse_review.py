"""Build the Synapse Review page: every logged synapse as a card to accept or reject from a phone.

Reads 'Synapse Log.xlsx' (one row per synapse, with the Accept and Notes typed or imported so far),
embeds the rows in 'synapse_review_page.html', and writes

  <Zarr Synapse ExF>\\synapse_review.html          open copy for this PC (images beside it)
  <zarr folder on the server>\\synapse_review.html  the shared copy, beside the gallery, locked with the
                                                    gallery password (same lock as Make_Gallery_Protected)

Decisions made in a browser stay in that browser until exported (Export CSV on the page); import the
CSV with review_import.py, run Measure_ExF.bat, then rebuild this page so the decisions are baked
into it for every device. The card images come from review_thumbs.py (run first; this script runs it
when the JPEG folder is missing).

Usage:  python synapse_review.py [--no-server] [--password]      (GALLERY_PASSWORD env var skips the prompt)
"""
import configparser
import datetime as dt
import html
import json
import re
import sys
import urllib.parse
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import convert_to_zarr as core       # noqa: E402  settings and share_url_for
import make_gallery                  # noqa: E402  lock() and ask_password()

ini = configparser.ConfigParser()
ini.read(HERE / "settings.ini", encoding="utf-8")
SHOTS = Path(ini.get("paths", "screenshot_folder"))
EXF = SHOTS.parent / "Zarr Synapse ExF"
LOG = EXF / "Synapse Log.xlsx"
TEMPLATE = HERE / "synapse_review_page.html"
TITLE = "Synapse Review"
PAGE_NAME = "synapse_review.html"
HEADER_ROW = 3
NG_PREFIX = "https://neuroglancer-demo.appspot.com/#!https://neuroglancer.kuanlab.yale.edu/data/exm/2026/zarr/_ng_states/synapses/"
GALLERY_REL = "../gallery/Zarr%20Synapse%20ExF/"   # from the server zarr folder to the mirrored synapse folder


def num(x):
    try:
        return None if x in (None, "", "–") else float(x)
    except (TypeError, ValueError):
        return None


def rows_from_log():
    """One compact dict per logged synapse, with the picture name, the Neuroglancer state and any
    Accept/Notes already in the log."""
    ws = load_workbook(LOG)["Synapses"]
    hdr = [c.value for c in ws[HEADER_ROW]]
    ci = {h: i for i, h in enumerate(hdr) if h}
    out = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        v = [c.value for c in row]
        if v[ci["Sample"]] is None:
            continue
        pic = row[ci["Picture"]].hyperlink
        ng = row[ci["Neuroglancer"]].hyperlink
        img = Path(pic.target).stem if pic else None
        if not img:
            continue
        ng_t = ng.target if ng else ""
        d = {"k": f"{v[ci['Sample']]}__{v[ci['Synapse ID']]}", "s": v[ci["Sample"]], "id": v[ci["Synapse ID"]], "n": int(v[ci["#"]]),
             "img": img, "ab": v[ci["Antibody"]], "abs": num(v[ci["AB Signal (σ)"]]), "abp": v[ci["AB+"]] == "Y",
             "ab2s": num(v[ci["AB2 Signal (σ)"]]), "ab2p": v[ci["AB2+"]] == "Y",
             "exf": num(v[ci["ExF (peak to peak)"]]), "exfg": num(v[ci["ExF (Gaussian fit)"]]), "in": v[ci["In 16–24"]] == "Y",
             "cls": v[ci["Class"]], "by": v[ci["PSD Named By"]], "score": num(v[ci["Example Score"]]), "r2": num(v[ci["Fit R²"]]),
             "dip": num(v[ci["Dip"]]), "pr": num(v[ci["Peak Ratio"]]), "z": v[ci["Slice"]],
             "ng": ng_t[len(NG_PREFIX):] if ng_t.startswith(NG_PREFIX) else None,
             "acc": (str(v[ci["Accept"]]).strip().upper() if v[ci["Accept"]] else None),
             "notes": (str(v[ci["Notes"]]).strip() if v[ci["Notes"]] else None),
             "rcls": (str(v[ci["Reviewed Class"]]).strip() if "Reviewed Class" in ci and v[ci["Reviewed Class"]] else None)}
        out.append({k: x for k, x in d.items() if x not in (None, False, "")} | {"abp": d["abp"], "ab2p": d["ab2p"], "in": d["in"]})
    return out


def with_manifest(rows):
    """Mouse, genotype and notes from datasets.csv (written by review_import.py) onto each row of that dataset."""
    import csv
    manifest = EXF / "datasets.csv"
    if not manifest.is_file():
        return rows
    with open(manifest, newline="", encoding="utf-8-sig") as f:
        meta = {r["dataset"]: r for r in csv.DictReader(f)}
    for r in rows:
        m = meta.get(r["s"])
        if m:
            for key, col in (("mouse", "animal"), ("geno", "genotype"), ("snotes", "notes")):
                if m.get(col):
                    r[key] = m[col]
    return rows


def build(rows, img_base, png_base):
    samples = sorted({r["s"] for r in rows})
    meta = {"built": f"{dt.datetime.now():%Y-%m-%d %H:%M}", "n": len(rows), "samples": samples,
            "img_base": img_base, "png_base": png_base, "ng_prefix": NG_PREFIX}
    data = json.dumps({"meta": meta, "rows": rows}, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    page = TEMPLATE.read_text(encoding="utf-8")
    return page.replace("__TITLE__", html.escape(TITLE)).replace("__BUILT__", html.escape(meta["built"])).replace("__DATA__", data)


def main():
    if not (EXF / "_review_jpg").is_dir():
        print("No card JPEGs yet: running review_thumbs.py first.")
        import review_thumbs
        review_thumbs.main()
    rows = with_manifest(rows_from_log())
    decided = sum(1 for r in rows if r.get("acc") in ("Y", "N"))
    print(f"{len(rows)} logged synapses from {len({r['s'] for r in rows})} samples; {decided} already accepted or rejected in the log.")

    local = EXF / PAGE_NAME
    local.write_text(build(rows, "_review_jpg/", ""), encoding="utf-8")
    print(f"PC copy: {local}")

    if "--no-server" not in sys.argv:
        password = make_gallery.ask_password() if ("--password" in sys.argv or __import__("os").environ.get("GALLERY_PASSWORD")) else None
        page = build(rows, GALLERY_REL + "_review_jpg/", GALLERY_REL)
        out = Path(core.DEFAULT_OUT) / PAGE_NAME
        out.write_text(make_gallery.lock(page, password, TITLE) if password else page, encoding="utf-8")
        share = core.share_url_for(out)
        print(f"Server copy: {out}  ({out.stat().st_size / 1e6:.1f} MB){'  PASSWORD PROTECTED' if password else '  (no password: pass --password)'}")
        print(f"Share: {share}")
        make_gallery.to_clipboard(share)


if __name__ == "__main__":
    main()
