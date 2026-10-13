"""
Size of the cell nuclei in expanded stacks, for an expansion factor from nuclei.

  nucleus_size.py [text]      go through the stacks whose names contain the text

M'Saad, Cairns et al. (Nat. Biotechnol. 2025) calibrated their expansion factor on nuclei:
ExF = square root of (mean nuclear cross-sectional area expanded / mean area not expanded),
with SYTOX Green-stained nuclei imaged in both conditions. That needs nuclei measured in tissue
that was NOT expanded. Put those numbers in PRE_EXPANSION below and the expansion factors
appear; until then this reports the expanded sizes, and what size they imply before expansion
at the expansion factor the synapses gave.

Nuclei need no dye here. In the pan stain a nucleus is a large, round, smooth region (little
fine structure) with a bright nucleolus. On each plane of a coarse pyramid level the smooth
regions are outlined; outlines on neighbouring planes are joined into nuclei. A nucleus counts
only when it is complete: it stays clear of the image border and of the first and last gel
planes, and its largest section has smaller ones on both sides.

Two sizes per nucleus, giving the two expansion factors:
  volume     sum of its section areas x the plane spacing (Cavalieri's estimator).
             ExF = cube root of (volume expanded / volume before)
  projection area of all its sections laid on top of each other (a maximum-intensity
             projection of the nucleus, as a wide-field image of a stained nucleus shows).
             ExF = square root of (area expanded / area before)
"""

import csv
import math
import sys
from pathlib import Path

import numpy as np

import ab_efficiency as eff
import convert_to_zarr as core
import roi_log as roi
import synapse_exf as syn

# Nuclei of the same tissue and region, NOT expanded: fill these in to get expansion factors.
PRE_EXPANSION = {
    "volume_um3": None,        # mean nuclear volume, µm³
    "projected_area_um2": None,  # mean projected (or equatorial cross-sectional) nuclear area, µm²
    "source": "",              # where the numbers come from, for the record
}
WORK_UM = 0.6            # nuclei are outlined on the pyramid level nearest this pixel size (as imaged)
DIAMETER_UM = (80.0, 400.0)   # a section this wide, as imaged, can be a nucleus (5-25 µm at 16x)
MIN_SOLIDITY = 0.85      # ...when it is this convex
MAX_ECCENTRICITY = 0.85  # ...and this round
SMOOTH_BELOW = 0.75      # smooth = local roughness under this share of Otsu's threshold
JOIN_UM = 40.0           # sections on neighbouring planes within this distance are one nucleus
MIN_SECTIONS = 3
CSV_NAME = "Nucleus Sizes.csv"
COLUMNS = ["Sample", "Complete Nuclei", "Volume (µm³, as imaged)", "Projected Area (µm², as imaged)",
           "ExF (nucleus volume)", "ExF (nucleus projection)", "Reference", "Synapse ExF",
           "Implied Diameter Before Expansion, from Volume (µm)", "Implied Diameter Before Expansion, from Projection (µm)",
           "Plane Spacing (µm, as imaged)", "Picture"]


def sections(plane, um):
    """Outlines of nuclear sections on one plane: [(mask, centre row, centre col)]."""
    from scipy import ndimage
    from skimage import filters, measure, morphology

    if (plane > 0).mean() < 0.5:
        return []
    fine = np.abs(plane - ndimage.gaussian_filter(plane, 1.5))          # fine structure: what a nucleus lacks
    rough = ndimage.gaussian_filter(ndimage.uniform_filter(fine, 15) / (ndimage.uniform_filter(plane, 15) + 1), 4)
    smooth = rough < SMOOTH_BELOW * filters.threshold_otsu(rough)
    smooth = morphology.remove_small_objects(morphology.binary_opening(smooth, morphology.disk(5)), 500)
    smooth = ndimage.binary_fill_holes(smooth)                         # the nucleolus is not smooth; it belongs
    found = []
    labels = measure.label(smooth)
    for r in measure.regionprops(labels):
        width = 2 * math.sqrt(r.area / math.pi) * um
        y0, x0, y1, x1 = r.bbox
        at_border = y0 == 0 or x0 == 0 or y1 == plane.shape[0] or x1 == plane.shape[1]
        # 'sure': large, round and convex enough to start a nucleus from. The smaller sections near
        # a nucleus's top and bottom fail that test; they are kept as well and only ever added to a
        # nucleus that has already been found, plane by plane outwards from it.
        sure = DIAMETER_UM[0] <= width <= DIAMETER_UM[1] and r.solidity >= MIN_SOLIDITY and r.eccentricity <= MAX_ECCENTRICITY
        if sure or (POLE_UM <= width <= DIAMETER_UM[1] and r.solidity >= 0.7):
            found.append({"mask": labels == r.label, "row": r.centroid[0], "col": r.centroid[1], "area": r.area * um * um,
                          "border": at_border, "sure": sure})
    return found


RENDER_UM = 2.4          # voxel size of the 3D rendering, as imaged
POLE_UM = 15.0           # the smallest section, as imaged, that can be the top or bottom of a nucleus
POLE_INSIDE = 0.6        # ...when this share of it lies over the section before it
MIN_DEPTH = 0.5          # a whole nucleus is at least this deep, as a share of its width


def render_3d(nucleus, um, z_um, png, name):
    """A 3D surface of one nucleus from its outlined sections, three views and a turning GIF.

    The planes are far apart (about 10 µm) compared with the pixels, so the surface between
    two sections is filled in by shape-based interpolation: each section becomes a map of
    signed distance to its outline, and the maps are blended linearly from plane to plane.
    The ends are closed one plane beyond the first and last section, shrunk by one plane
    spacing. The orange rings are the sections themselves: only those are measured.
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from PIL import Image
    from scipy import ndimage
    from skimage import measure

    sections = nucleus["sections"]
    rows, cols = np.nonzero(nucleus["shadow"])
    pad = int(round(12 / um))
    ya, yb = max(0, rows.min() - pad), rows.max() + pad
    xa, xb = max(0, cols.min() - pad), cols.max() + pad
    shrink = max(1, int(round(RENDER_UM / um)))
    step = um * shrink
    signed = {}
    for sec in sections:
        m = sec["mask"][ya:yb:shrink, xa:xb:shrink]
        signed[sec["z"]] = (ndimage.distance_transform_edt(m) - ndimage.distance_transform_edt(~m)) * step
    zs = sorted(signed)
    planes = np.arange(zs[0] - 1, zs[-1] + 2)
    known = np.array([zs[0] - 1] + zs + [zs[-1] + 1])
    maps = np.stack([signed[zs[0]] - z_um] + [signed[z] for z in zs] + [signed[zs[-1]] - z_um])
    # blend linearly from plane to plane, onto planes RENDER_UM apart
    fine = np.arange(planes[0], planes[-1] + 1e-6, step / z_um)
    upper = np.clip(np.searchsorted(known, fine, side="right"), 1, len(known) - 1)
    share = ((fine - known[upper - 1]) / (known[upper] - known[upper - 1]))[:, None, None]
    volume = maps[upper - 1] * (1 - share) + maps[upper] * share
    # Close the two ends (a layer of "outside" all round) and take the blockiness of the coarse
    # voxels off the surface. The smoothing is about one voxel: it rounds off bumps and pulls the
    # surface a little inside the orange rings. It is for looking at; the volume is not taken from it.
    volume = ndimage.gaussian_filter(np.pad(volume, 2, constant_values=-3 * z_um), 1.2)
    verts, faces, _, _ = measure.marching_cubes(volume, 0, spacing=(step, step, step))
    verts = verts - 2 * step
    z0 = (zs[0] - 1 - fine[0]) * z_um                        # depth of the first map, for placing the rings
    rings = []
    for sec in sections:
        depth = (sec["z"] - planes[0]) * z_um
        for line in measure.find_contours(sec["mask"][ya:yb, xa:xb].astype(float), 0.5):
            rings.append((line[:, 1] * um, line[:, 0] * um, np.full(len(line), depth)))
    size = np.ptp(verts, axis=0)                             # extent in z, y, x

    def draw(ax, elev, azim, title=None):
        ax.plot_trisurf(verts[:, 2], verts[:, 1], faces, verts[:, 0], color="#C8CDD4", shade=True, linewidth=0, antialiased=True)
        for x, y, z in rings:
            ax.plot(x, y, z, color="#FF9900", lw=0.9)
        ax.set_box_aspect((size[2], size[1], size[0]))
        ax.view_init(elev=elev, azim=azim)
        ax.set_facecolor(eff.SURFACE)
        for axis, label in ((ax.xaxis, "x (µm)"), (ax.yaxis, "y (µm)"), (ax.zaxis, "depth (µm)")):
            axis.set_pane_color((0.09, 0.09, 0.09, 1))
            axis.label.set_color(eff.INK)
            axis.set_label_text(label, fontsize=8)
            axis._axinfo["grid"]["color"] = (0.2, 0.2, 0.2, 1)
        ax.tick_params(colors=eff.MUTED, labelsize=7)
        ax.invert_zaxis()                                    # first plane on top, as in the stack
        if title:
            ax.set_title(title, color=eff.INK, fontsize=10)

    fig = plt.figure(figsize=(15, 5.6), facecolor=eff.SURFACE)
    for i, (elev, azim, title) in enumerate(((22, -60, "Tilted"), (90, -90, "From Above (XY)"), (0, -90, "From the Side (XZ)")), start=1):
        draw(fig.add_subplot(1, 3, i, projection="3d"), elev, azim, title)
    fig.suptitle(f"{name}: Nucleus at Slice {nucleus['equator']['z']} in 3D (as imaged)   ·   Volume {nucleus['volume']:,.0f} µm³,  "
                 f"{size[2]:.0f} × {size[1]:.0f} × {size[0]:.0f} µm (x × y × depth)\n"
                 f"Orange rings = the {len(sections)} measured sections, {z_um:.1f} µm apart; the surface between them is interpolated",
                 color=eff.INK, fontsize=10)
    eff.save(fig, png)

    # the same nucleus turning once, as a GIF beside the picture
    fig = plt.figure(figsize=(5, 5), facecolor=eff.SURFACE, dpi=90)
    ax = fig.add_subplot(1, 1, 1, projection="3d")
    frames = []
    for azim in range(0, 360, 10):
        ax.clear()
        draw(ax, 22, azim)
        fig.canvas.draw()
        frames.append(Image.fromarray(np.asarray(fig.canvas.buffer_rgba())[..., :3].copy()))
    plt.close(fig)
    gif = png.with_suffix(".gif")
    frames[0].save(gif, save_all=True, append_images=frames[1:], duration=110, loop=0)
    print(f"   {gif.name}")


def measure_sample(name, zarr_dir, out_root, folder):
    """Complete nuclei of one stack. Returns a row for the CSV, or raises ValueError."""
    series = zarr_dir / "0"
    info = core.read_series(series)
    if "stack" not in roi.classify(info):
        raise ValueError("not a stack: a single plane cuts nuclei at random heights")
    channels, pan, _, _ = eff.channel_roles(name, zarr_dir, out_root, info)
    if pan is None:
        raise ValueError("no pan stain found")
    full_um = eff.voxel_um(info, (1, info["shape"]["y"], info["shape"]["x"]))[2]
    level = int(np.clip(round(math.log2(WORK_UM / full_um)), 0, len(info["levels"]) - 1))
    raw = eff.open_level(series, info, level)[pan]                     # lazy (z, y, x)
    z_um, _, um = eff.voxel_um(info, raw.shape)
    if not z_um:
        raise ValueError("no plane spacing in the file")
    z0, z1 = eff.gel_planes(eff.open_level(series, info, eff.macro_level(info))[pan])
    nuclei = []                                                        # each a list of sections, in plane order
    every = {}
    for z in range(z0, z1 + 1):
        every[z] = sections(raw[z].compute().astype(np.float32), um)
        for s in every[z]:
            s["z"] = z
        for s in [s for s in every[z] if s["sure"]]:
            for n in nuclei:
                last = n[-1]
                if 0 < z - last["z"] <= 2 and math.hypot(s["row"] - last["row"], s["col"] - last["col"]) * um < JOIN_UM:
                    n.append(s)
                    break
            else:
                nuclei.append([s])
    # Follow each nucleus to its top and bottom: on the next plane out, the section that lies over
    # the last one and is no larger belongs to it. Without this the poles are cut off and the
    # nucleus comes out as a slab.
    for n in nuclei:
        for step in (-1, 1):
            edge = n[0] if step < 0 else n[-1]
            z = edge["z"] + step
            while z0 <= z <= z1:
                over = [s for s in every[z] if (s["mask"] & edge["mask"]).sum() >= POLE_INSIDE * s["mask"].sum()
                        and s["area"] <= 1.15 * edge["area"] and not any(s is t for t in n)]
                if not over:
                    break
                edge = max(over, key=lambda s: s["area"])
                n.insert(0, edge) if step < 0 else n.append(edge)
                z += step
    whole = []
    for n in nuclei:
        areas = [s["area"] for s in n]
        k = int(np.argmax(areas))
        width = 2 * math.sqrt(max(areas) / math.pi)
        if (len(n) >= MIN_SECTIONS and 0 < k < len(n) - 1 and not any(s["border"] for s in n)
                and n[0]["z"] > z0 and n[-1]["z"] < z1 and (n[-1]["z"] - n[0]["z"] + 1) * z_um >= MIN_DEPTH * width):
            # planes it skipped (a section not recognised) count with the mean of their neighbours
            zs = np.arange(n[0]["z"], n[-1]["z"] + 1)
            volume = float(np.interp(zs, [s["z"] for s in n], areas).sum() * z_um)
            shadow = np.any([s["mask"] for s in n], axis=0)
            whole.append({"volume": volume, "projected": float(shadow.sum() * um * um), "shadow": shadow, "equator": n[k],
                          "sections": n})
    print(f"   {len(nuclei)} candidate(s), {len(whole)} complete; level {level} ({um:.2f} µm pixels), planes {z0}-{z1}, spacing {z_um:.1f} µm")
    if not whole:
        raise ValueError("no complete nucleus in the stack")

    volume = np.array([w["volume"] for w in whole])
    area = np.array([w["projected"] for w in whole])
    spread = lambda v: f"{v.mean():.0f} ± {v.std(ddof=1) if len(v) > 1 else 0:.0f}"
    row = {"Sample": name, "Complete Nuclei": len(whole), "Volume (µm³, as imaged)": spread(volume),
           "Projected Area (µm², as imaged)": spread(area), "Plane Spacing (µm, as imaged)": round(z_um, 2),
           "Reference": PRE_EXPANSION["source"] or "none: no nuclei measured before expansion"}
    if PRE_EXPANSION["volume_um3"]:
        row["ExF (nucleus volume)"] = round(float((volume.mean() / PRE_EXPANSION["volume_um3"]) ** (1 / 3)), 1)
    if PRE_EXPANSION["projected_area_um2"]:
        row["ExF (nucleus projection)"] = round(float(math.sqrt(area.mean() / PRE_EXPANSION["projected_area_um2"])), 1)

    # one picture: each complete nucleus on its equatorial plane, with the outline of its projection
    from skimage import measure

    fig, axes = eff.figure(1, len(whole), figsize=(4.2 * len(whole), 4.6), squeeze=False)
    for ax, w in zip(axes[0], whole):
        plane = raw[w["equator"]["z"]].compute().astype(np.float32)
        r, c = int(w["equator"]["row"]), int(w["equator"]["col"])
        half = int(0.75 * DIAMETER_UM[1] / um / 2)
        ya, xa = max(0, r - half), max(0, c - half)
        crop = plane[ya:r + half, xa:c + half]
        lo, hi = np.percentile(crop, [1, 99.7])
        ax.imshow(np.clip((crop - lo) / max(hi - lo, 1), 0, 1), cmap="gray", interpolation="nearest")
        for line in measure.find_contours(w["shadow"][ya:r + half, xa:c + half].astype(float), 0.5):
            ax.plot(line[:, 1], line[:, 0], color="#FF9900", lw=1.5)
        bar = 50 / um
        ax.plot([6, 6 + bar], [crop.shape[0] - 8] * 2, color="white", lw=3, solid_capstyle="butt")
        ax.text(6, crop.shape[0] - 12, "50 µm", color="white", fontsize=8, va="bottom")
        ax.set_title(f"Slice {w['equator']['z']}:  Volume {w['volume']:,.0f} µm³,  Projection {w['projected']:,.0f} µm²",
                     color=eff.INK, fontsize=9)
        ax.set_xticks([])
        ax.set_yticks([])
    fig.suptitle(f"{name}: Complete Nuclei in the Pan Stain (as imaged). Orange = outline of the nucleus projected through its planes",
                 color=eff.INK, fontsize=10)
    fig.tight_layout()
    png = folder / f"{name}_nuclei.png"
    eff.save(fig, png)
    row["Picture"] = str(png)
    # the largest complete nucleus in 3D
    render_3d(max(whole, key=lambda w: w["volume"]), um, z_um, folder / f"{name}_nuclei-3d.png", name)
    return row, volume.mean(), area.mean()


def synapse_exf_of(base):
    """{sample: ExF from its synapses} from the Synapse Log, where it is trustworthy."""
    from openpyxl import load_workbook

    try:
        lines = load_workbook(base / syn.LOG_NAME, read_only=True)["Summary"].iter_rows(values_only=True)
        heads = next(lines)
        rows = [dict(zip(heads, line)) for line in lines]
    except Exception:
        return {}
    return {r["Sample"]: r["ExF"] for r in rows if r.get("ExF") and not str(r.get("ExF Basis") or "").startswith("all double peaks")}


def run(text):
    core.load_settings()
    out_root = Path(core.DEFAULT_OUT)
    base = Path(core.DEFAULT_SHOTS).parent / "Zarr Synapse ExF"
    base.mkdir(parents=True, exist_ok=True)
    from_synapses = synapse_exf_of(base)
    kept = {}
    path = base / CSV_NAME
    if path.is_file():
        with open(path, newline="", encoding="utf-8-sig") as f:
            kept = {r["Sample"]: r for r in csv.DictReader(f)}
    for name, zarr_dir in roi.samples(out_root):
        if text.lower() not in name.lower() or syn.ignored(name):
            continue
        print(name)
        folder = base / name
        try:
            folder.mkdir(exist_ok=True)
            row, volume, area = measure_sample(name, zarr_dir, out_root, folder)
        except Exception as e:
            print(f"   skipped: {e}")
            continue
        exf = from_synapses.get(name)
        if exf:
            row["Synapse ExF"] = exf
            row["Implied Diameter Before Expansion, from Volume (µm)"] = round((6 * volume / math.pi) ** (1 / 3) / exf, 1)
            row["Implied Diameter Before Expansion, from Projection (µm)"] = round(2 * math.sqrt(area / math.pi) / exf, 1)
        kept[name] = row
        print("   " + "; ".join(f"{k}: {v}" for k, v in row.items() if k not in ("Sample", "Picture")))
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(kept[k] for k in sorted(kept))
    print(f"\n{path}")
    if not PRE_EXPANSION["volume_um3"] and not PRE_EXPANSION["projected_area_um2"]:
        print("No expansion factor from nuclei: PRE_EXPANSION at the top of nucleus_size.py is empty.\n"
              "Measure nuclei in tissue that was not expanded, put the mean volume and/or projected area there, and run again.")


if __name__ == "__main__":
    run(sys.argv[1] if len(sys.argv) > 1 else "")
