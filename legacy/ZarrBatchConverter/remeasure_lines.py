"""Re-measure synapses along lines placed by hand on the Synapse Review page ('Fix line').

The page saves two points per synapse in plane pixels: the dense-projection (DP) end first, then the
PSD end. review_import.py writes them into the Synapse Log's 'Manual Line' column as 'x1 y1 x2 y2'.
This script reads the pan stain along that line exactly as synapse_exf.py reads it along its own
(same plane, Gaussian smoothing of SMOOTH_PX, a lane LINE_PX pixels thick, samples every half pixel,
each top placed between samples by a parabola through its brightest sample and two neighbours,
searched within a third of the line's length of each end) and writes

  P2P Manual (µm)    peak-to-peak distance along the hand-placed line, as imaged
  ExF (manual line)  that distance / 81.9 nm

into the log. A check picture per synapse (the tile with the line, and the profile with both tops)
goes to <Zarr Synapse ExF>/<sample>/manual_lines/, and 'Manual vs automatic.csv' beside the log
pairs each manual reading with the automatic one on the same synapse, with the Bland-Altman summary
(mean difference and 95 % limits of agreement), the hand-measurement check of the analysis plan (1.3).

Usage:  python remeasure_lines.py [--all] [--log PATH]
        --all   re-measure every line, not only those without a result yet
"""
import argparse
import configparser
import csv
import math
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import convert_to_zarr as core          # noqa: E402  settings (ZBC_SETTINGS-aware) and zarr reading
import ab_efficiency as eff             # noqa: E402  channel roles, levels, voxel size, figure style
import synapse_exf as sx                # noqa: E402  the measurement constants

HEADER_ROW = 3
LINE, P2P, EXF_COL = "Manual Line", "P2P Manual (µm)", "ExF (manual line)"


def log_path():
    core.load_settings()
    cfg = configparser.ConfigParser(interpolation=None)
    cfg.read(core.SETTINGS_FILE, encoding="utf-8-sig")
    folder = cfg.get("paths", "synapse_folder", fallback="").strip() or str(Path(core.DEFAULT_SHOTS).parent / "Zarr Synapse ExF")
    return Path(folder) / "Synapse Log.xlsx"


def parse_line(text):
    try:
        v = [float(x) for x in str(text or "").replace(",", " ").split()]
    except ValueError:
        return None
    return v if len(v) == 4 else None


def profile_along(plane, a, b, um):
    """The pan stain along the line from a (DP end) to b (PSD end), both (x, y) in plane pixels.
    Returns (t, thin, summits, distance_px, edge) with t in pixels from a, or None if the line runs
    off the plane."""
    from scipy import ndimage

    a, b = np.asarray(a, float), np.asarray(b, float)
    length = float(np.hypot(*(b - a)))
    if length < 2:
        return None
    u = (b - a) / length                        # (x, y) along the line
    n = np.array([-u[1], u[0]])                 # across it
    t = np.arange(-0.5 * length, 1.5 * length + 0.25, 0.5)
    lane = np.arange(sx.LINE_PX) - (sx.LINE_PX - 1) / 2
    xs = a[0] + t[:, None] * u[0] + lane[None, :] * n[0]
    ys = a[1] + t[:, None] * u[1] + lane[None, :] * n[1]
    if xs.min() < 1 or ys.min() < 1 or xs.max() > plane.shape[1] - 2 or ys.max() > plane.shape[0] - 2:
        return None
    thin = ndimage.map_coordinates(plane, [ys, xs], order=1).mean(axis=1)
    edge = False

    def summit(m):
        nonlocal edge
        near = np.flatnonzero(np.abs(t - m) <= length / 3)
        i = near[np.argmax(thin[near])]
        if i in (near[0], near[-1]):
            edge = True                             # the brightest sample is at the end of its window
        if not 0 < i < len(t) - 1:
            return float(t[i])
        bend = thin[i - 1] - 2 * thin[i] + thin[i + 1]
        return float(t[i] + float(np.clip(0.5 * (thin[i - 1] - thin[i + 1]) / bend if bend < 0 else 0.0, -0.5, 0.5)) * (t[1] - t[0]))

    tops = (summit(0.0), summit(length))
    return t, thin, tops, abs(tops[1] - tops[0]), edge


def check_picture(png, plane, a, b, res, um, auto_exf, manual_exf, title):
    t, thin, tops, _, edge = res
    half = int(max(40, 1.2 * np.hypot(b[0] - a[0], b[1] - a[1])))
    cx, cy = int(round((a[0] + b[0]) / 2)), int(round((a[1] + b[1]) / 2))
    tile = plane[max(0, cy - half):cy + half + 1, max(0, cx - half):cx + half + 1]
    ox, oy = max(0, cx - half), max(0, cy - half)
    fig, axes = eff.figure(1, 2, figsize=(9, 4))
    lo, hi = np.percentile(tile, [1, 99.8])
    axes[0].imshow(np.clip((tile - lo) / max(hi - lo, 1e-6), 0, 1), cmap="gray", interpolation="nearest")
    axes[0].plot([a[0] - ox, b[0] - ox], [a[1] - oy, b[1] - oy], color="#2fd6e8", lw=2)
    axes[0].plot([a[0] - ox], [a[1] - oy], "o", color=sx.DP_LIME, ms=5)
    axes[0].plot([b[0] - ox], [b[1] - oy], "o", color=sx.PSD_PINK, ms=5)
    axes[0].set_xticks([]); axes[0].set_yticks([])
    axes[0].set_title("Your line (DP lime, PSD salmon)", color=eff.INK, fontsize=9)
    axes[1].plot(t * um, thin, color=sx.MEASURED_RED, lw=2)
    for k, m in enumerate(tops):
        axes[1].axvline(m * um, color=sx.DP_LIME if k == 0 else sx.PSD_PINK, lw=1, ls=":")
    axes[1].set_xlabel("Distance along your line from the DP end (µm, as imaged)", color=eff.INK, fontsize=9)
    axes[1].set_ylabel("Counts", color=eff.INK, fontsize=9)
    axes[1].set_title(f"ExF along your line {manual_exf:.1f}   (automatic {auto_exf if auto_exf is not None else '–'})"
                      + ("   · a top sits at the edge of its window" if edge else ""), color=eff.INK, fontsize=9, loc="left")
    fig.suptitle(title, color=eff.INK, fontsize=9, x=0.01, ha="left")
    png.parent.mkdir(parents=True, exist_ok=True)
    eff.save(fig, png, 150)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--all", action="store_true", help="re-measure every line, not only new ones")
    ap.add_argument("--log", help="a Synapse Log other than the one in the settings")
    ap.add_argument("--no-pictures", action="store_true")
    args = ap.parse_args()
    log = Path(args.log) if args.log else log_path()
    core.load_settings()
    zroot = Path(core.DEFAULT_OUT)
    try:
        wb = load_workbook(log)
    except PermissionError:
        sys.exit(f"{log.name} is open in Excel. Close it and run this again.")
    ws = wb["Synapses"]
    cols = {c.value: c.column for c in ws[HEADER_ROW] if c.value}
    if LINE not in cols:
        sys.exit("The log has no 'Manual Line' column yet: place lines on the review page and import them first.")
    todo = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        cell = lambda name: row[cols[name] - 1]
        line = parse_line(cell(LINE).value) if cell("Sample").value else None
        if line and (args.all or cell(P2P).value in (None, "")):
            todo.append((row, line))
    print(f"{len(todo)} hand-placed line(s) to measure in {log}")
    if not todo:
        return 0

    mine = PatternFill("solid", fgColor="FFF2CC")
    by_sample = {}
    for row, line in todo:
        by_sample.setdefault(row[cols["Sample"] - 1].value, []).append((row, line))
    pairs = []
    from scipy import ndimage
    for name, items in by_sample.items():
        zdir = zroot / f"{name}.zarr"
        try:
            info = core.read_series(zdir / "0")
            channels, pan, _, _ = eff.channel_roles(name, zdir, zroot, info)
            raw = eff.open_level(zdir / "0", info, 0)
            um = eff.voxel_um(info, raw.shape[1:])[2]
        except Exception as e:
            print(f"   {name}: cannot read the image ({e}); skipped")
            continue
        if pan is None:
            print(f"   {name}: no pan channel; skipped")
            continue
        for row, (x1, y1, x2, y2) in items:
            z = int(row[cols["Slice"] - 1].value)
            pad = int(np.hypot(x2 - x1, y2 - y1) * 1.6) + 12
            c0, r0 = int(max(0, min(x1, x2) - pad)), int(max(0, min(y1, y2) - pad))
            block = raw[pan, z, r0:int(max(y1, y2) + pad) + 1, c0:int(max(x1, x2) + pad) + 1].compute().astype(np.float32)
            plane = ndimage.gaussian_filter(block, sx.SMOOTH_PX)
            a, b = (x1 - c0, y1 - r0), (x2 - c0, y2 - r0)
            res = profile_along(plane, a, b, um)
            n = row[cols["#"] - 1].value
            if res is None:
                print(f"   {name} #{n}: the line runs off the image; skipped")
                continue
            p2p_um = res[3] * um
            exf = p2p_um * 1000 / sx.REF_NM
            for col, value in ((P2P, round(p2p_um, 3)), (EXF_COL, round(exf, 1))):
                c = row[cols[col] - 1]; c.value, c.fill = value, mine
            auto = row[cols["ExF (peak to peak)"] - 1].value
            accept = row[cols["Accept"] - 1].value
            pairs.append({"Sample": name, "#": n, "Accept": accept or "", "ExF automatic": auto, "ExF manual line": round(exf, 2),
                          "Difference (manual - automatic)": None if auto is None else round(exf - float(auto), 2),
                          "Top at window edge": "Y" if res[4] else ""})
            print(f"   {name} #{n}: ExF {exf:.1f} along your line (automatic {auto})" + ("   [a top at the window edge]" if res[4] else ""))
            if not args.no_pictures:
                pic = row[cols["Picture"] - 1].hyperlink
                stem = Path(pic.target).stem if pic else f"{name}_synapse-{int(n):03d}"
                check_picture(log.parent / name / "manual_lines" / f"{stem}_manual.png", plane, a, b, res, um,
                              auto, exf, f"{name}  #{n}  slice {z}")
    try:
        wb.save(log)
    except PermissionError:
        sys.exit(f"{log.name} is open in Excel; nothing saved. Close it and run this again.")
    print(f"saved {log}")

    # agreement with the automatic reading, over every synapse that has both (not only this run's)
    every = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        m, a = row[cols[EXF_COL] - 1].value, row[cols["ExF (peak to peak)"] - 1].value
        if m not in (None, "") and a not in (None, ""):
            every.append((row[cols["Sample"] - 1].value, row[cols["#"] - 1].value, row[cols["Accept"] - 1].value or "", float(a), float(m)))
    out = log.parent / "Manual vs automatic.csv"
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Sample", "#", "Accept", "ExF automatic", "ExF manual line", "Difference (manual - automatic)"])
        for s, n, acc, a, m in every:
            w.writerow([s, n, acc, a, m, round(m - a, 2)])
        if len(every) >= 3:
            d = np.array([m - a for *_, a, m in every])
            mean, sd = float(d.mean()), float(d.std(ddof=1))
            w.writerow([]); w.writerow(["Bland-Altman", f"n = {len(d)}", "", "mean difference", round(mean, 2), f"95% limits {mean - 1.96 * sd:.2f} to {mean + 1.96 * sd:.2f}"])
            print(f"\nManual line versus automatic, n = {len(d)}: mean difference {mean:+.2f} ExF, 95 % limits of agreement {mean - 1.96 * sd:+.2f} to {mean + 1.96 * sd:+.2f}")
    print(f"agreement table: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
