"""
Build the gallery page ('2026 AC Gallery.html') in the screenshot folder: one page, opened in a browser, for
looking through every sample quickly.

  - a grid of thumbnails, one card per sample
  - search, filter by type, sort by name / date / antibody density / rating
  - click a card for all of that sample's pictures, its metadata, and its Neuroglancer link

It is built from what is already on disk - the pictures in each sample's subfolder, the
screenshot log (ratings, notes), the image list (metadata, links) and the synapse log
(expansion factor, resolution, antibody labelling; see synapse_exf.py) - so it needs no new
exports and can be rebuilt at any time. Samples or pictures that do not exist yet are simply
left out. Small JPEG copies of the pictures are kept in '_gallery_thumbs' so the page loads
fast; the full-size PNG opens when a picture is clicked.
"""

import datetime as dt
import html
import json
import re
import sys
import urllib.parse
from pathlib import Path

from openpyxl import load_workbook

import convert_to_zarr as core
import roi_log

GALLERY_TITLE = "2026 AC Gallery"   # shown at the top of the page; also the file name
CURRENT_YEAR = "2026"                # the year the page opens on; older years come from settings_<year>.ini
GALLERY_NAME = GALLERY_TITLE + ".html"
OLD_NAMES = ["Zarr Gallery.html"]    # earlier names of the page, removed when it is rebuilt
THUMB_DIR = "_gallery_thumbs"
THUMB_PX = 420
SYNAPSE_THUMB_PX = 1100   # synapse panels are wide strips, so their small copies are wider
SYNAPSE_FOLDER = "Zarr Synapse ExF"   # beside the screenshot folder; written by Measure_ExF.bat
# the order pictures appear in a sample's page; the first one found is the card's picture
VIEW_ORDER = ["Volume", "MIP-XY", "MIP-XZ", "MIP-YZ", "Overview", "Full-image"] + \
             [f"{kind}-{q:02d}pct" for q in roi_log.PLANE_PERCENTS for kind in ("Plane", "Zoom-200pct")] + \
             ["Zoom-200pct", "Seam-crop", "Crop-centre", "Crop-top-left", "Crop-top-right",
              "Crop-bottom-left", "Crop-bottom-right"]
CARD_ORDER = ["Volume", "Plane-50pct", "Overview", "Full-image", "MIP-XY"]


def nice_view(tag):
    """'Zoom-200pct-25pct' -> '200% close-up at 25% depth'."""
    m = re.fullmatch(r"Plane-(\d+)pct", tag)
    if m:
        return f"Plane at {int(m.group(1))}% depth"
    m = re.fullmatch(r"Zoom-(\d+)pct-(\d+)pct", tag)
    if m:
        return f"{m.group(1)}% close-up at {int(m.group(2))}% depth"
    m = re.fullmatch(r"Zoom-(\d+)pct", tag)
    if m:
        return f"{m.group(1)}% close-up"
    return {"MIP-XY": "Projection, top view (XY)", "MIP-XZ": "Projection, side view (XZ)",
            "MIP-YZ": "Projection, side view (YZ)", "Volume": "Tilted volume",
            "Full-image": "Full image"}.get(tag, tag.replace("-", " "))


def public_json(samples):
    """The samples as the page's data, without the local folder each was read from."""
    clean = [{k: v for k, v in s.items() if not k.startswith("_")} for s in samples]
    return json.dumps(clean, ensure_ascii=False, default=str).replace("</", "<\\/")


def thumbnail(png, dest, size=THUMB_PX):
    """Make (or reuse) the small JPEG copy of a picture. Returns False if it cannot be read."""
    from PIL import Image

    try:
        if dest.is_file() and dest.stat().st_mtime >= png.stat().st_mtime:
            return True
        img = Image.open(png).convert("RGB")
        img.thumbnail((size, size))
        dest.parent.mkdir(parents=True, exist_ok=True)
        img.save(dest, "JPEG", quality=82)
        return True
    except Exception:
        return False   # still being written, or not a picture


def read_metadata(image_list=None, shots=None):
    """{image name: metadata} from the image list, and {name: [log rows]} from the screenshot log
    (the current year's unless another year's list and screenshot folder are given)."""
    meta = {}
    try:
        ws = load_workbook(image_list or core.DEFAULT_LIST)["Images"]
        header_row = next(r for r in range(1, 12) if "Image Name" in [c.value for c in ws[r]])
        heads = [c.value for c in ws[header_row]]
        for r in range(header_row + 1, ws.max_row + 1):
            row = {h: ws.cell(r, i + 1).value for i, h in enumerate(heads) if h}
            link = ws.cell(r, heads.index("Image Name") + 1).hyperlink
            if row.get("Image Name"):
                row["url"] = None if link is None else (link.target or "") + ("#" + link.location if link.location else "")
                meta[row["Image Name"]] = row
    except Exception as e:
        print(f"(image list not read: {e})")
    logged = {}
    try:
        for row in roi_log.read_log(Path(shots or core.DEFAULT_SHOTS) / roi_log.LOG_NAME):
            logged.setdefault(row.get(roi_log.HEADERS[0]), []).append(row)
    except Exception as e:
        print(f"(screenshot log not read: {e})")
    return meta, logged


SYNAPSE_LINKS = {}   # (sample, synapse number) -> Neuroglancer link, from the Synapse Log; filled by collect()


def read_synapse_links(folder):
    """The Neuroglancer link of each logged synapse, as the Synapse Log has it. The log is
    written in the same run as the pictures, so a number means the same synapse in both."""
    links = {}
    try:
        ws = load_workbook(folder / "Synapse Log.xlsx")["Synapses"]
        heads = [c.value for c in ws[3]]
        at = {h: heads.index(h) for h in ("#", "Sample", "Neuroglancer")}
        for row in ws.iter_rows(min_row=4):
            link = row[at["Neuroglancer"]].hyperlink
            if link is not None and row[at["Sample"]].value:
                links[(row[at["Sample"]].value, int(row[at["#"]].value))] = link.target
    except Exception as e:
        print(f"(synapse links not read: {e})")
    return links


HIDDEN = ("tetraspec", "pff")   # samples left out of the gallery: bead images, and the PFF samples


def synapse_pictures(base, name):
    """The synapse panels of one sample, in the order they are shown: the contact sheet
    first, then each synapse by its number in the Synapse Log. Empty if none were made."""
    folder = base.parent / SYNAPSE_FOLDER / name
    if not folder.is_dir():
        return []
    found = []
    for png in folder.glob("*.png"):
        figure = png.name == f"{name}_synapses_structure.png"      # the structure measurements, shown with the sheet
        solid = png.name == f"{name}_nuclei-3d.png"                # the largest nucleus rendered in 3D
        nuclei = solid or png.name == f"{name}_nuclei.png"         # the nuclei measured by nucleus_size.py
        sheet = figure or nuclei or png.name == f"{name}_synapses_contact-sheet.png"
        m = re.fullmatch(re.escape(name) + r"_synapse-(\d+)_z(\d+)(?:_x\d+_y\d+)?\.png", png.name)
        if not sheet and not m:
            continue
        thumb = base / THUMB_DIR / "_synapses" / name / (png.stem + ".jpg")
        if not thumbnail(png, thumb, SYNAPSE_THUMB_PX):
            continue
        found.append({
            "tag": "synapse", "group": "synapse", "number": 0 if sheet else int(m.group(1)),
            "view": "Nucleus in 3D" if solid else "Nuclei (for an ExF from nuclei)" if nuclei else "Structure measurements" if figure else "All synapses (contact sheet)" if sheet else f"Synapse #{int(m.group(1))}",
            "slice": None if sheet else int(m.group(2)),
            "file": urllib.parse.quote(f"../{SYNAPSE_FOLDER}/{name}/{png.name}"), "filename": png.name,
            "thumb": urllib.parse.quote(f"{THUMB_DIR}/_synapses/{name}/{thumb.name}"),
            "rating": None, "notes": None,
            # clicking a synapse's thumbnail opens Neuroglancer on it, when the log has a link for it
            "ng": SYNAPSE_LINKS.get((name, int(m.group(1)))) if m else None,
        })
    return sorted(found, key=lambda p: p["number"])


def read_synapse_summary(base, folder=None):
    """{sample: its row of the Synapse Log's Summary sheet} (synapse_exf.py), and the folder the
    synapse pictures are in. Empty when nothing has been measured yet."""
    import configparser

    if folder is None:
        cfg = configparser.ConfigParser(interpolation=None)
        cfg.read(core.SETTINGS_FILE, encoding="utf-8-sig")
        folder = cfg.get("paths", "synapse_folder", fallback="").strip()
    folder = Path(folder or base.parent / "Zarr Synapse ExF")
    rows = {}
    try:
        ws = load_workbook(folder / "Synapse Log.xlsx", read_only=True)["Summary"]
        lines = ws.iter_rows(values_only=True)
        heads = next(lines)
        for line in lines:
            row = dict(zip(heads, line))
            if row.get("Sample"):
                rows[row["Sample"]] = row
    except Exception as e:
        print(f"(synapse log not read: {e})")
    return rows, folder


def read_channels(base):
    """{sample: [its channels]} from the channel check (channel_check.py): what each channel is,
    its colour, the antibody on it when that is known, and whether it shows any signal."""
    import csv

    found = {}
    try:
        with open(base.parent / "Zarr Comparisons" / "Channel Check - Channels.csv", newline="", encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                found.setdefault(r["Sample"], []).append({
                    "role": r.get("Role") or "antibody", "laser": r.get("Laser (nm)") or "", "antibody": r.get("Antibody") or "",
                    "colour": r.get("Colour") or "#ffffff", "dark": (r.get("Status") or "").startswith("DARK"),
                    "status": r.get("Status") or "", "channel": r.get("Channel") or ""})
    except Exception as e:
        print(f"(channel check not read: {e})")
    return found


def labelling(name, channels):
    """How a sample is labelled, for the tabs and filters: its chips, how many antibody channels
    show signal, the combination as words, and the genotype read from the name."""
    chips, shown = [], []
    for c in sorted(channels, key=lambda c: (c["role"] != "pan", c["role"] != "antibody", -int(c["laser"] or 0))):
        if c["role"] == "pan":
            text = f"pan · {c['laser']}"
        elif c["role"] == "antibody":
            text = f"{c['antibody'] or 'antibody ?'} · {c['laser']}"
            if not c["dark"]:
                shown.append(c["antibody"] or f"{c['laser']} nm")
        else:
            text = f"nuclear · {c['laser']}"
        chips.append({"text": text, "colour": c["colour"], "dark": c["dark"], "role": c["role"],
                      "title": c["channel"] + ": " + (c["status"] or "no check yet")})
    has_pan = any(c["role"] == "pan" for c in channels)
    kind = {1: "single", 2: "double"}.get(len(shown), "triple" if len(shown) >= 3 else ("pan" if has_pan else ""))
    genotype = next((t[:2].upper() for t in name.split("_") if re.fullmatch(r"(?i)(wt|ko)\d*[a-z]?", t)), "")
    return {"chips": chips, "label_kind": kind, "combo": " + ".join(sorted(shown, key=str.lower)), "genotype": genotype}


def older_years():
    """[(year, {paths})] for every settings_<year>.ini beside this script (made when the
    pre-2026 stacks were converted), newest first. Each year keeps its own zarrs, image list,
    screenshots and synapse log on the server."""
    import configparser

    years = []
    for ini in sorted(Path(__file__).parent.glob("settings_[0-9][0-9][0-9][0-9].ini"), reverse=True):
        cfg = configparser.ConfigParser(interpolation=None)
        cfg.read(ini, encoding="utf-8-sig")
        paths = {k: cfg.get("paths", k, fallback="").strip() for k in ("screenshot_folder", "image_list", "synapse_folder")}
        if paths["screenshot_folder"] and Path(paths["screenshot_folder"]).is_dir():
            years.append((ini.stem.split("_")[1], paths))
    return years


def collect(base, year=None, image_list=None, synapse_folder=None, root=""):
    """One record per sample subfolder that holds at least one picture. For another year, root
    is the web address of its screenshot folder: its pictures are not beside the page."""
    meta, logged = read_metadata(image_list, base if image_list else None)
    measured, synapse_folder = read_synapse_summary(base, synapse_folder)
    checked = read_channels(base)
    SYNAPSE_LINKS.update(read_synapse_links(synapse_folder))
    samples = []
    folders = sorted((d for d in base.iterdir() if d.is_dir() and not d.name.startswith("_")
                      and not any(word in d.name.lower() for word in HIDDEN)),
                     key=lambda d: d.name.lower())
    for n, folder in enumerate(folders, start=1):
        name = folder.name
        notes = {r.get("Screenshot"): r for r in logged.get(name, [])}
        pictures = []
        for png in folder.glob("*.png"):
            m = re.fullmatch(re.escape(name) + r"_(.+)_z(all|\d+)(?:_\d+)?\.png", png.name)
            if not m:
                continue
            tag = m.group(1)
            thumb = base / THUMB_DIR / name / (png.stem + ".jpg")
            if not thumbnail(png, thumb):
                continue
            row = notes.get(png.name, {})
            pictures.append({
                "tag": tag, "view": nice_view(tag), "slice": None if m.group(2) == "all" else int(m.group(2)),
                "file": urllib.parse.quote(f"{name}/{png.name}"),
                "thumb": urllib.parse.quote(f"{THUMB_DIR}/{name}/{thumb.name}"),
                "rating": row.get("Rating"), "notes": row.get("Notes"),
            })
        if not pictures:
            continue
        pictures.sort(key=lambda p: (VIEW_ORDER.index(p["tag"]) if p["tag"] in VIEW_ORDER else len(VIEW_ORDER),
                                     p["tag"], p["slice"] or 0))
        x = measured.get(name, {})   # this sample's row of the synapse log's Summary sheet
        card = next((p for tag in CARD_ORDER for p in pictures if p["tag"] == tag), pictures[0])
        synapses = synapse_pictures(base, name)
        pictures = pictures + synapses   # always after every other picture of the sample
        m = meta.get(name, {})
        ratings = [p["rating"] for p in pictures if isinstance(p["rating"], (int, float))]
        date = m.get("Date Acquired")
        if root:   # another year: every address made absolute, so the page can sit anywhere
            pictures = [dict(p, file=root + p["file"], thumb=root + p["thumb"]) for p in pictures]
            card = dict(card, thumb=root + card["thumb"])
        samples.append({
            **labelling(name, checked.get(name, [])),
            "year": year or "", "root": root, "_base": str(base),
            "name": name, "card": card["thumb"], "pictures": pictures, "synapses": len([p for p in synapses if p["number"]]),
            "type": next((r.get("Type") for r in logged.get(name, []) if r.get("Type")), ""),
            "date": date.strftime("%Y-%m-%d") if hasattr(date, "strftime") else (str(date) if date else ""),
            "channels": m.get("Channels (laser)") or "", "microscope": m.get("Microscope") or "",
            "objective": m.get("Objective") or "", "size": m.get("Size X × Y × Z") or "",
            "extent": m.get("Extent X × Y × Z (µm)") or "", "density": m.get("AB Spots / 1000 µm²"),
            "density_text": m.get("AB Density by Channel") or "", "link_status": m.get("Link Status") or "",
            "url": m.get("url"), "rating": max(ratings) if ratings else None,
            "notes": " | ".join(str(p["notes"]) for p in pictures if p["notes"]),
            "folder": root + urllib.parse.quote(name) + "/",
            # from the synapse log; absent until Measure_ExF.bat has been run on the sample
            "antibody": x.get("Antibody") or "", "exf": x.get("ExF"), "exf_sd": x.get("ExF s.d."),
            "exf_fields": x.get("ExF Fields"), "exf_n": x.get("ExF n"), "exf_basis": x.get("ExF Basis") or "",
            "exf_fit": x.get("ExF (Gaussian fit)"), "exf_ci": x.get("ExF 95% CI") or "",
            # an expansion factor from pan-only double peaks reads high and is shown as doubtful
            "exf_sure": bool(x.get("ExF")) and not str(x.get("ExF Basis") or "").startswith("all double peaks"),
            "pixel": x.get("Pixel (nm, corrected)"), "pixel_raw": x.get("Pixel (nm, as imaged)"),
            "res": x.get("Resolution (nm, corrected)"), "res_raw": x.get("Resolution FWHM (nm, as imaged)"),
            "synapses_found": x.get("Synapses found"), "ab_pos": x.get("AB+"), "ab_labelled": x.get("AB Labelled Synapses (%)"),
            "ab_signal": x.get("AB Signal at Synapses (σ)"), "ab_puncta": x.get("AB Puncta"),
            "ab_on": x.get("AB Puncta on Synapses (%)"), "ab_enrich": x.get("AB Enrichment (× chance)"),
            "synapse_folder": root + urllib.parse.quote(f"../{synapse_folder.name}/{name}") + "/" if x else "",
            "type1_macular": x.get("Type I Macular"), "type1_perforated": x.get("Type I Perforated"),
            "type2": x.get("Type II (putative)"), "classes_on": x.get("Classes Counted On") or "",
            # structure measured on the same synapses (synapse_structure.py)
            "dpdp": x.get("DP–DP (nm, corrected)"), "dpdp_n": x.get("DP–DP n"), "exf_dpdp": x.get("ExF from DP–DP (÷67 nm)"),
            "aniso": x.get("ExF Anisotropy (%)"), "aniso_p": x.get("Anisotropy p"), "field_cv": x.get("ExF Field CV (%)"),
            "ab_psd": x.get("AB to PSD (nm, corrected; + = towards DP)"), "ab_dp": x.get("AB to DP (nm, corrected)"),
            "psd_len": x.get("PSD Length (nm, corrected)"), "psd_thick": x.get("PSD Thickness (nm, corrected, with blur)"),
            "lab_faces": x.get("AB Labelled Near Faces (%)"), "lab_centre": x.get("AB Labelled At Centre (%)"),
            "confirmed": x.get("Confirmed Synapses / 100 µm² (corrected)"), "perforated": x.get("Perforated (%)"),
        })
        if n % 20 == 0:
            print(f"   {n}/{len(folders)} samples")
    return samples


PAGE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITLE__</title>
<style>
:root { --bg:#111418; --panel:#1b2027; --line:#2c333d; --text:#e6e9ee; --dim:#9aa4b2; --accent:#4fc3b0; --warn:#ff8a80; }
* { box-sizing:border-box; }
body { margin:0; background:var(--bg); color:var(--text); font:14px/1.45 "Segoe UI",system-ui,sans-serif; }
header { position:sticky; top:0; z-index:5; background:rgba(17,20,24,.96); border-bottom:1px solid var(--line);
         padding:12px 18px; display:flex; flex-wrap:wrap; gap:10px; align-items:center; }
h1 { font-size:18px; margin:0 14px 0 0; font-weight:600; }
input,select { background:var(--panel); color:var(--text); border:1px solid var(--line); border-radius:6px; padding:7px 9px; font:inherit; }
input[type=search] { width:260px; }
#count { color:var(--dim); margin-left:auto; }
#grid { display:grid; grid-template-columns:repeat(auto-fill,minmax(230px,1fr)); gap:14px; padding:18px; }
.card { background:var(--panel); border:1px solid var(--line); border-radius:8px; overflow:hidden; cursor:pointer; }
.card:hover,.card:focus { border-color:var(--accent); outline:none; }
.card img { width:100%; aspect-ratio:1/1; object-fit:contain; background:#000; display:block; }
.card .t { padding:8px 10px 10px; }
.name { font-weight:600; font-size:12.5px; word-break:break-word; }
.sub { color:var(--dim); font-size:12px; margin-top:3px; }
.badge { display:inline-block; font-size:11px; padding:1px 7px; border-radius:10px; border:1px solid var(--line); color:var(--dim); margin:4px 4px 0 0; }
.badge.r { color:#ffd54f; border-color:#6b5a1e; } .badge.w { color:var(--warn); border-color:#6b2a25; }
a.pic { color:var(--accent); text-decoration:none; } a.pic:hover { text-decoration:underline; }
.fn { font:11px/1.35 Consolas,monospace; color:var(--dim); user-select:all; word-break:break-all; cursor:text; }
.badge.x { color:var(--accent); border-color:#27695f; } .badge.q { border-style:dashed; }
.chips { margin-top:5px; display:flex; flex-wrap:wrap; gap:4px 6px; }
.chip { display:inline-flex; align-items:center; gap:5px; font-size:11.5px; padding:1px 8px 1px 6px; border-radius:10px;
        border:1px solid var(--line); color:var(--text); white-space:nowrap; }
.chip i { width:9px; height:9px; border-radius:50%; background:var(--c); border:1px solid var(--c); display:inline-block; }
.chip.dark { color:var(--dim); border-style:dashed; } .chip.dark i { background:transparent; }
.chips.big .chip { font-size:13px; padding:3px 10px 3px 8px; } .chips.big { margin:8px 0 2px; }
#tabs { display:flex; flex-wrap:wrap; gap:6px; padding:10px 18px 0; }
#tabs button { background:none; border:1px solid var(--line); color:var(--dim); border-radius:16px; padding:5px 13px; font:inherit; cursor:pointer; }
#tabs button.on { background:var(--panel); color:var(--text); border-color:var(--accent); }
#tabs button span { color:var(--dim); margin-left:5px; font-size:12px; }
.group.yr { color:var(--text); font-size:14px; }
.group { grid-column:1 / -1; margin:8px 0 -4px; font-size:13px; font-weight:600; color:var(--dim); border-bottom:1px solid var(--line); padding-bottom:5px; }
#detail { position:fixed; inset:0; background:rgba(0,0,0,.82); display:none; z-index:10; overflow:auto; }
#detail.on { display:block; }
#box { background:var(--bg); max-width:1500px; margin:22px auto; border:1px solid var(--line); border-radius:10px; padding:18px 22px 26px; }
#box h2 { margin:0 0 4px; font-size:17px; word-break:break-word; }
.bar { display:flex; gap:10px; align-items:center; flex-wrap:wrap; margin:10px 0 14px; }
.btn { background:var(--panel); border:1px solid var(--line); color:var(--text); border-radius:6px; padding:6px 12px; text-decoration:none; cursor:pointer; font:inherit; }
.btn:hover { border-color:var(--accent); } .btn.go { background:#17443d; border-color:#27695f; }
table { border-collapse:collapse; margin:4px 0 16px; } td { padding:2px 14px 2px 0; vertical-align:top; } td:first-child { color:var(--dim); white-space:nowrap; }
#pics { display:grid; grid-template-columns:repeat(auto-fill,minmax(250px,1fr)); gap:14px; }
figure { margin:0; background:var(--panel); border:1px solid var(--line); border-radius:8px; overflow:hidden; }
figure img { width:100%; aspect-ratio:1/1; object-fit:contain; background:#000; display:block; }
figcaption { padding:7px 10px 9px; font-size:12.5px; } figcaption span { color:var(--dim); }
.empty { padding:60px; text-align:center; color:var(--dim); }
h3 { margin:26px 0 10px; font-size:14px; font-weight:600; color:var(--dim); border-top:1px solid var(--line); padding-top:16px; }
#syn { display:grid; grid-template-columns:repeat(auto-fill,minmax(620px,1fr)); gap:14px; }
#syn figure img { aspect-ratio:auto; }
</style></head><body>
<header>
  <h1>__TITLE__</h1>
  <select id="year" title="Which year's images to show"></select>
  <input id="q" type="search" placeholder="Search name, channels, notes…" autofocus>
  <select id="type"><option value="">All types</option></select>
  <select id="objective"><option value="">All objectives</option></select>
  <select id="genotype"><option value="">All genotypes</option></select>
  <select id="sort">
    <option value="name">Sort: name</option><option value="date">Sort: date, newest first</option>
    <option value="density">Sort: antibody density</option><option value="rating">Sort: rating</option>
    <option value="exf">Sort: expansion factor</option><option value="res">Sort: resolution, finest first</option>
    <option value="enrich">Sort: antibody on synapses (enrichment)</option>
    <option value="labelled">Sort: synapses labelled (%)</option>
  </select>
  <label><input id="rated" type="checkbox"> rated only</label>
  <label><input id="measured" type="checkbox"> with expansion factor only</label>
  <span id="count"></span>
</header>
<nav id="tabs"></nav>
<div id="grid"></div>
<div id="detail"><div id="box"></div></div>
<script>
const DATA = __DATA__;
const BUILT = "__BUILT__";
const $ = id => document.getElementById(id);
const esc = s => String(s ?? "").replace(/[&<>"]/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
[...new Set(DATA.map(s => s.type).filter(Boolean))].sort().forEach(t => $("type").insertAdjacentHTML("beforeend", `<option>${esc(t)}</option>`));
[...new Set(DATA.map(s => s.objective).filter(Boolean))].sort().forEach(t => $("objective").insertAdjacentHTML("beforeend", `<option>${esc(t)}</option>`));
[...new Set(DATA.map(s => s.genotype).filter(Boolean))].sort().forEach(t => $("genotype").insertAdjacentHTML("beforeend", `<option>${esc(t)}</option>`));
// Year: the page opens on the newest year; 'before' is every older year, 'all' every year with the
// newest first. Kept in the address (?year=2024) like the tab.
const YEARS = [...new Set(DATA.map(s => s.year).filter(Boolean))].sort().reverse();
const NEWEST = YEARS[0] || "";
let YEAR = new URLSearchParams(location.search).get("year") || NEWEST;
const yearOptions = [...YEARS.map(y => [y, y]), ...(YEARS.length > 2 ? [["before", "Before " + NEWEST]] : []), ...(YEARS.length > 1 ? [["all", "All years"]] : [])];
if (!yearOptions.some(o => o[0] === YEAR)) YEAR = NEWEST;
$("year").innerHTML = yearOptions.map(([v, text]) => `<option value="${v}">${text} · ${v === "all" ? DATA.length : DATA.filter(s => v === "before" ? s.year !== NEWEST : s.year === v).length}</option>`).join("");
$("year").value = YEAR;
$("year").hidden = YEARS.length < 2;
$("year").addEventListener("input", () => { YEAR = $("year").value; address(); tabs(); draw(); });
const inYear = s => YEAR === "all" || (YEAR === "before" ? s.year !== NEWEST : s.year === YEAR);
const address = () => { const q = new URLSearchParams(); if (YEAR !== NEWEST) q.set("year", YEAR); if (TAB) q.set("tab", TAB);
  history.replaceState(null, "", location.pathname + (q.toString() ? "?" + q : "")); };
// Tabs: how many antibody channels show signal. The tab is kept in the address (?tab=double) so it can be shared.
const TABS = [["", "All"], ["pan", "Pan only"], ["single", "Single label"], ["double", "Double label"], ["triple", "Triple label"]];
let TAB = new URLSearchParams(location.search).get("tab") || "";
if (!TABS.some(t => t[0] === TAB)) TAB = "";
function tabs() {
  $("tabs").innerHTML = TABS.map(([key, text]) => {
    const n = DATA.filter(s => inYear(s) && (!key || s.label_kind === key)).length;
    return n || !key ? `<button data-tab="${key}" class="${key === TAB ? "on" : ""}">${text}<span>${n}</span></button>` : "";
  }).join("");
}
$("tabs").addEventListener("click", e => {
  const b = e.target.closest("button"); if (!b) return;
  TAB = b.dataset.tab;
  address(); tabs(); draw();
});
// One chip per channel, a dot in the channel's Neuroglancer colour; hollow and dashed when the channel shows no signal.
const chips = s => (s.chips || []).map(c => `<span class="chip${c.dark ? " dark" : ""}" style="--c:${esc(c.colour)}" title="${esc(c.title)}${c.dark ? " (no signal)" : ""}"><i></i>${esc(c.text)}</span>`).join("");
const grouped = () => (TAB === "double" || TAB === "triple") && $("sort").value === "name";
let shown = [];
function draw() {
  const q = $("q").value.trim().toLowerCase(), type = $("type").value, sort = $("sort").value, rated = $("rated").checked;
  const measured = $("measured").checked, objective = $("objective").value, genotype = $("genotype").value;
  shown = DATA.filter(s => inYear(s) && (!TAB || s.label_kind === TAB) && (!objective || s.objective === objective) && (!genotype || s.genotype === genotype) &&
    (!type || s.type === type) && (!rated || s.rating != null) && (!measured || s.exf != null) &&
    (!q || q.split(/\s+/).every(w => (s.name + " " + s.channels + " " + s.notes + " " + s.type + " " + s.date + " " + s.antibody + " " + (s.combo || "") + " " + (s.genotype || "") + " " + (s.chips || []).map(c => c.text).join(" ")).toLowerCase().includes(w))));
  const by = { name:(a,b)=>a.name.localeCompare(b.name), date:(a,b)=>b.date.localeCompare(a.date),
               density:(a,b)=>(b.density ?? -1)-(a.density ?? -1), rating:(a,b)=>(b.rating ?? -1)-(a.rating ?? -1),
               exf:(a,b)=>(b.exf ?? -1)-(a.exf ?? -1), res:(a,b)=>(a.res ?? 1e9)-(b.res ?? 1e9),
               enrich:(a,b)=>(b.ab_enrich ?? -1)-(a.ab_enrich ?? -1), labelled:(a,b)=>(b.ab_labelled ?? -1)-(a.ab_labelled ?? -1) }[sort];
  shown.sort(by);
  if (grouped()) shown.sort((a, b) => (a.combo || "").localeCompare(b.combo || "") || a.name.localeCompare(b.name));   // by combination of antibodies
  const several = new Set(shown.map(s => s.year)).size > 1;
  if (several) shown.sort((a, b) => (b.year || "").localeCompare(a.year || ""));   // stable: the newest year first, each sorted as chosen
  const heading = (s, i) => (several && (i === 0 || shown[i - 1].year !== s.year)
      ? `<div class="group yr">${esc(s.year || "year not known")} · ${shown.filter(x => x.year === s.year).length}</div>` : "") +
    (grouped() && (i === 0 || shown[i - 1].combo !== s.combo || shown[i - 1].year !== s.year)
      ? `<div class="group">${esc(s.combo || "combination not known")} · ${shown.filter(x => x.combo === s.combo && x.year === s.year).length}</div>` : "");
  $("count").textContent = `${shown.length} of ${DATA.filter(inYear).length} samples · built ${BUILT}`;
  $("grid").innerHTML = shown.length ? shown.map((s,i) => `${heading(s, i)}
    <div class="card" tabindex="0" data-i="${i}">
      <img loading="lazy" src="${s.card}" alt="">
      <div class="t"><div class="name">${esc(s.name)}</div>
        <div class="sub">${esc(s.date)}${s.type ? " · " + esc(s.type) : ""} · ${s.pictures.length - (s.synapses ? s.synapses + 1 : 0)} images</div>
        <div class="chips">${chips(s)}</div>
        ${s.rating != null ? `<span class="badge r">★ ${s.rating}</span>` : ""}
        ${s.density != null ? `<span class="badge">${Number(s.density).toFixed(1)} spots / 1000 µm²</span>` : ""}
        ${s.exf != null ? `<span class="badge x${s.exf_sure ? "" : " q"}" title="${esc(s.exf_basis)}">ExF ${s.exf_sure ? "=" : "≈"} ${Number(s.exf).toFixed(1)}</span>` : ""}
        ${s.res != null ? `<span class="badge" title="width of the smallest antibody puncta ÷ ExF: an upper bound">≤ ${Number(s.res).toFixed(0)} nm</span>` : ""}
        ${s.ab_enrich != null ? `<span class="badge" title="antibody puncta on synapses, × chance">AB ${Number(s.ab_enrich).toFixed(1)}× on synapses</span>` : ""}
        ${s.synapses ? `<span class="badge">${s.synapses} synapses</span>` : ""}
        ${s.link_status && s.link_status !== "OK" ? `<span class="badge w">link problem</span>` : ""}
      </div></div>`).join("") : `<div class="empty">No samples match.</div>`;
}
const fig = p => `<figure><a href="${p.ng ? esc(p.ng) : p.file}" target="_blank" rel="noopener"
    title="${p.ng ? "Open this synapse in Neuroglancer" : "Open the full-size image"}"><img loading="lazy" src="${p.thumb}" alt=""></a>
  <figcaption>${esc(p.view)}${p.slice != null ? ` <span>· slice ${p.slice}</span>` : ""}
  ${p.ng ? ` <span>·</span> <a class="pic" href="${p.file}" target="_blank">image</a>` : ""}
  ${p.filename ? `<br><code class="fn" title="click to select, then Ctrl+C">${esc(p.filename)}</code>` : ""}
  ${p.rating != null ? ` <span>· ★ ${p.rating}</span>` : ""}${p.notes ? `<br><span>${esc(p.notes)}</span>` : ""}</figcaption></figure>`;
function open_(i) {
  const s = shown[i]; if (!s) return;
  const rows = [["Date", s.date],["Type", s.type],["Channels", s.channels],["Microscope", s.microscope],["Objective", s.objective],
    ["Size (voxels)", s.size],["Extent (µm)", s.extent],["Antibody density", s.density_text],["Link status", s.link_status],["Notes", s.notes],
    ["Expansion factor", s.exf == null ? "" : `ExF ${s.exf_sure ? "=" : "≈"} ${s.exf} ± ${s.exf_sd}${s.exf_ci ? `, 95% CI ${s.exf_ci}` : ""}  (DP–PSD peak to peak ÷ 81.9 nm; ${s.exf_basis}; ` +
      `${s.exf_fields ? s.exf_fields + " fields, " : ""}${s.exf_n} synapses; Gaussian fit: ExF = ${s.exf_fit})` +
      (s.exf_sure ? "" : "  — not reliable: no synaptic antibody to confirm the synapses, reads high")],
    ["Pixel size", s.pixel == null ? "" : `${s.pixel} nm corrected  (${s.pixel_raw} nm as imaged)`],
    ["Resolution", s.res == null ? "" : `≤ ${s.res} nm corrected  (${s.res_raw} nm as imaged; smallest antibody puncta, lateral, an upper bound)`],
    ["Synapses found", s.synapses_found == null ? "" : `${s.synapses_found}${s.antibody ? `, ${s.ab_pos} with ${s.antibody}` : ""}`],
    ["Synapse classes", s.type1_macular == null ? "" : `Type I macular ${s.type1_macular}, Type I perforated ${s.type1_perforated}, ` +
      `Type II (putative) ${s.type2}  (counted on ${s.classes_on}; automatic, after M'Saad, Cairns et al. Fig. 2)`],
    ["Expansion checks", s.aniso == null && s.exf_dpdp == null ? "" :
      (s.aniso != null ? `by orientation: ${s.aniso}% amplitude, p = ${s.aniso_p}` : "") +
      (s.field_cv != null ? `; between fields: CV ${s.field_cv}%` : "") +
      (s.exf_dpdp != null ? `; second ruler, DP–DP ${s.dpdp} nm (n = ${s.dpdp_n}) ÷ 67 nm: ExF = ${s.exf_dpdp}` : "")],
    ["PSD", s.psd_len == null ? "" : `length ${s.psd_len} nm, thickness ${s.psd_thick} nm including optical blur` +
      (s.perforated != null ? `; ${s.perforated}% perforated` : "") + "  (corrected, mean ± s.d.)"],
    ["Antibody position", s.ab_psd == null ? "" : `${s.ab_psd} nm from the PSD peak (+ = towards the DPs), ${s.ab_dp} nm from the DP peak  (corrected, mean ± s.d.)`],
    ["Labelling by depth", s.lab_faces == null ? "" : `${s.lab_faces}% of synapses labelled near the gel faces, ${s.lab_centre}% at the centre`],
    ["Confirmed synapses", s.confirmed == null ? "" : `${s.confirmed} per 100 µm² corrected  (antibody-positive, side-on only: a lower bound)`],
    ["Antibody labelling", s.ab_labelled == null ? "" : `${s.ab_labelled}% of synapses found carry antibody` +
      (s.ab_signal != null ? ` (median ${s.ab_signal} σ)` : "") +
      (s.ab_on != null ? `; ${s.ab_on}% of ${s.ab_puncta} puncta sit on a synapse found = ${s.ab_enrich}× chance` : "") + "  (provisional)"]]
    .filter(r => r[1]).map(r => `<tr><td>${r[0]}</td><td>${esc(r[1])}</td></tr>`).join("");
  $("box").innerHTML = `<h2>${esc(s.name)}</h2>
    <div class="chips big">${chips(s)}</div>
    <div class="bar">
      ${s.url ? `<a class="btn go" href="${esc(s.url)}" target="_blank" rel="noopener">Open in Neuroglancer</a>` : ""}
      <a class="btn" href="${s.folder}" target="_blank">Open folder</a>
      ${s.synapse_folder ? `<a class="btn" href="${s.synapse_folder}" target="_blank">Synapse images</a>` : ""}
      ${s.review ? `<a class="btn go" href="${s.review}" title="Accept or reject this sample's synapses (phone friendly)">Review synapses</a>` : ""}
      <button class="btn" data-go="${i-1}">◀ Previous</button><button class="btn" data-go="${i+1}">Next ▶</button>
      <button class="btn" data-go="x">Close (Esc)</button>
    </div>
    <table>${rows}</table>
    <div id="pics">${s.pictures.filter(p => p.group !== "synapse").map(fig).join("")}</div>
    ${s.pictures.some(p => p.group === "synapse") ? `<h3>Synapses · ${s.synapses} measured for expansion factor</h3>
      <div id="syn">${s.pictures.filter(p => p.group === "synapse").map(fig).join("")}</div>` : ""}`;
  $("detail").classList.add("on"); $("detail").scrollTop = 0; $("detail").dataset.i = i;
  history.replaceState(null, "", "#" + encodeURIComponent(s.name));   // the address now points at this sample
}
const close_ = () => { $("detail").classList.remove("on"); history.replaceState(null, "", location.pathname + location.search); };
$("grid").addEventListener("click", e => { const c = e.target.closest(".card"); if (c) open_(+c.dataset.i); });
$("grid").addEventListener("keydown", e => { if (e.key === "Enter") { const c = e.target.closest(".card"); if (c) open_(+c.dataset.i); } });
$("detail").addEventListener("click", e => {
  if (e.target.id === "detail") return close_();
  const b = e.target.closest("[data-go]"); if (!b) return;
  b.dataset.go === "x" ? close_() : open_(Math.max(0, Math.min(shown.length - 1, +b.dataset.go)));
});
document.addEventListener("keydown", e => {
  if (!$("detail").classList.contains("on")) return;
  const i = +$("detail").dataset.i;
  if (e.key === "Escape") close_(); else if (e.key === "ArrowRight") open_(Math.min(shown.length - 1, i + 1)); else if (e.key === "ArrowLeft") open_(Math.max(0, i - 1));
});
["q","type","objective","genotype","sort","rated","measured"].forEach(id => $(id).addEventListener("input", draw));
tabs();
draw();
if (location.hash) {   // opened with a sample's name after '#': go straight to it
  const want = decodeURIComponent(location.hash.slice(1)), at = shown.findIndex(s => s.name === want);
  if (at >= 0) open_(at);
}
</script></body></html>
"""


SERVER_THUMB_PX = 300     # in the server copy: an ordinary picture
SERVER_SHEET_PX = 1000    # ...a contact sheet, structure figure or nuclei picture
SERVER_SYNAPSE_PX = 360   # ...one synapse's panel
# What the server copy needs on top of the ordinary page: nothing on it links to a file on
# this PC, so links with nowhere to go are hidden and do nothing.
SERVER_EXTRA = """<style>a.btn[href=""], a.pic[href=""] { display:none; } a[href=""] { cursor:default; }
#servernote { background:#2a2410; border-bottom:1px solid #6b5a1e; color:#f0e2b0; padding:8px 18px; font-size:13px; line-height:1.5; }
#servernote b { color:#ffd54f; } #servernote summary { cursor:pointer; }</style>
<script>
document.addEventListener("click", e => { if (e.target.closest('a[href=""]')) e.preventDefault(); }, true);
document.querySelector("header").insertAdjacentHTML("afterend", `<details id="servernote"><summary><b>Server copy.</b>
  Small images only; a few things from the full gallery are not here. Click for the list.</summary>
  <b>Works here:</b> search, filters and sorting; every sample's details (expansion factor, resolution, synapse classes, antibody labelling);
  "Open in Neuroglancer"; clicking a synapse to open it in Neuroglancer.<br>
  <b>Not here:</b> full-size images (these are reduced copies: ordinary images 300 pixels, contact sheets and figures 1000, single synapses 360);
  the "Open folder" and "Synapse images" buttons; the turning 3D nucleus GIFs; the Synapse Log workbook, the per-sample CSV files and the reports.<br>
  <b>Full version:</b> __FULL__ on the PC it was built on. This copy is rebuilt by Make_Gallery.bat and can be behind it; the build time is in the top bar.
</details>`);
</script>
</body>"""


def server_copy(base, samples, password=None):
    """The gallery as ONE file in the zarr folder on the lab server, so that anyone who can reach
    the server can open it, from the drive or from its web address, with nothing else installed.

    Every picture is carried inside the page as a small JPEG; the full-size pictures stay on
    this PC and are not linked. The Neuroglancer links work as usual, since they already
    point at the server. With a password the file holds only the encrypted page (see lock).
    Returns (path, address to share).
    """
    import base64
    import io

    from PIL import Image

    def packed(relative, px, quality, sample=None):
        root, local = (sample or {}).get("root", ""), Path((sample or {}).get("_base") or base)
        if root and relative.startswith(root):
            relative = relative[len(root):]
        try:
            img = Image.open(local / urllib.parse.unquote(relative)).convert("RGB")
        except Exception:
            return ""
        img.thumbnail((px, px))
        buffer = io.BytesIO()
        img.save(buffer, "JPEG", quality=quality, optimize=True)
        return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")

    carried = []
    for n, sample in enumerate(samples, start=1):
        sample = dict(sample, folder="", synapse_folder="",
                      review=f"synapse_review.html?sample={urllib.parse.quote(sample['name'])}"
                      if sample.get("synapses") and not sample.get("root") else "")
        pictures = []
        for picture in sample["pictures"]:
            synapse = picture.get("group") == "synapse"
            px, quality = ((SERVER_SYNAPSE_PX, 55) if picture.get("number") else (SERVER_SHEET_PX, 65)) if synapse else (SERVER_THUMB_PX, 70)
            pictures.append(dict(picture, thumb=packed(picture["thumb"], px, quality, sample), file=""))
        sample["pictures"] = pictures
        sample["card"] = packed(sample["card"], SERVER_THUMB_PX, 70, sample)
        carried.append(sample)
        if n % 20 == 0:
            print(f"   {n}/{len(samples)} samples packed")
    data = public_json(carried)
    title = GALLERY_TITLE + " (server copy)"
    page = PAGE.replace("__DATA__", data).replace("__BUILT__", html.escape(__import__("provenance").stamp(Path(__file__))))
    page = page.replace("__TITLE__", html.escape(title)).replace(
        "</body>", SERVER_EXTRA.replace("__FULL__", html.escape(str(base / GALLERY_NAME)).replace("\\", "\\\\")))
    out = Path(core.DEFAULT_OUT) / GALLERY_NAME
    out.write_text(lock(page, password, GALLERY_TITLE) if password else page, encoding="utf-8")
    return out, core.share_url_for(out)


MIRROR_FOLDER = "gallery"   # beside the zarr folder on the server: the full copy of the gallery's folders


def mirror(base):
    """Keep a full copy of the gallery on the lab server: the screenshot folder, the synapse
    folder and the comparisons, laid out as on this PC, so the page there has everything
    (full-size pictures, synapse pictures, reports) at a web address.

    It is a mirror: what is new or changed here is copied, and what was deleted here is
    deleted there. The working copies stay on this PC. Returns the page's web address.
    """
    import subprocess

    target = Path(core.DEFAULT_OUT).parent / MIRROR_FOLDER
    for folder in (base, base.parent / SYNAPSE_FOLDER, base.parent / "Zarr Comparisons"):
        if not folder.is_dir():
            continue
        # robocopy: /MIR mirror, /MT threads, /R /W few retries (a file open in Excel is skipped),
        # /XF Excel's lock files, /XD the measurement caches; exit codes under 8 mean success
        done = subprocess.run(["robocopy", str(folder), str(target / folder.name), "/MIR", "/MT:16", "/R:1", "/W:1",
                               "/NFL", "/NDL", "/NP", "/NJH", "/XF", "~$*", "/XD", "_measured"], capture_output=True, text=True)
        tail = [line.strip() for line in done.stdout.splitlines() if line.strip().startswith(("Files", "Bytes"))]
        print(f"   {folder.name}: " + ("; ".join(" ".join(t.split()) for t in tail) if done.returncode < 8 else f"robocopy failed ({done.returncode})"))
    return core.share_url_for(target / base.name / GALLERY_NAME)


# --------------------------------------------------------------------------- sharing

PBKDF2_ROUNDS = 600_000      # key stretching for the password: slows guessing, about half a second in a browser

LOCKED_PAGE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex, nofollow"><title>__TITLE__</title>
<style>
:root { color-scheme: light dark; --bg:#111418; --text:#e6e9ee; --dim:#9aa4b2; --line:#2c333d; --accent:#6db3ff; }
@media (prefers-color-scheme: light) { :root { --bg:#ffffff; --text:#1b1f24; --dim:#5b6470; --line:#d9dee5; --accent:#0b5cad; } }
body { margin:0; min-height:100vh; display:grid; place-items:center; background:var(--bg); color:var(--text); font:15px/1.5 Arial, sans-serif; }
main { width:min(360px, calc(100vw - 32px)); }
h1 { font-size:20px; margin:0 0 6px; } p { color:var(--dim); margin:0 0 16px; }
input, button { width:100%; box-sizing:border-box; font:inherit; padding:10px 12px; border-radius:6px; border:1px solid var(--line); }
input { background:transparent; color:var(--text); margin-bottom:10px; }
button { background:var(--accent); color:#fff; border:none; cursor:pointer; }
#m { margin-top:12px; min-height:1.5em; }
</style></head><body><main>
<h1>__TITLE__</h1><p>This page is password protected.</p>
<form id="f"><input id="p" type="password" autocomplete="current-password" placeholder="Password" autofocus required><button>Open</button></form>
<p id="m"></p></main>
<script>
(() => {   // its own scope: names declared at the top level would survive document.open and clash with the gallery's own
const SALT = "__SALT__", IV = "__IV__", ROUNDS = __ROUNDS__, DATA = "__DATA__";
const REMEMBER = "gallery-unlock", HOURS = 12;   // the password stays on this device for a while, so the gallery and the review page open each other
const bytes = async s => new Uint8Array(await (await fetch("data:application/octet-stream;base64," + s)).arrayBuffer());
const m = document.getElementById("m");
async function unlock(password, typed) {
  if (!window.crypto || !crypto.subtle) { m.textContent = "This browser cannot open protected pages here (it needs https)."; return false; }
  m.textContent = "Opening...";
  try {
    const base = await crypto.subtle.importKey("raw", new TextEncoder().encode(password), "PBKDF2", false, ["deriveKey"]);
    const key = await crypto.subtle.deriveKey({name: "PBKDF2", salt: await bytes(SALT), iterations: ROUNDS, hash: "SHA-256"},
                                              base, {name: "AES-GCM", length: 256}, false, ["decrypt"]);
    const page = await crypto.subtle.decrypt({name: "AES-GCM", iv: await bytes(IV)}, key, await bytes(DATA));
    try { localStorage.setItem(REMEMBER, JSON.stringify({p: password, until: Date.now() + HOURS * 3600e3})); } catch (e) {}
    document.open(); document.write(new TextDecoder().decode(page)); document.close();
    return true;
  } catch (err) { m.textContent = typed ? "Wrong password." : ""; return false; }
}
document.getElementById("f").addEventListener("submit", e => { e.preventDefault(); unlock(document.getElementById("p").value, true); });
try {   // a password remembered by the gallery or the review page within the last HOURS opens this page too
  const r = JSON.parse(localStorage.getItem(REMEMBER) || "null");
  if (r && r.until > Date.now()) unlock(r.p, false); else localStorage.removeItem(REMEMBER);
} catch (e) {}
})();
</script></body></html>"""


def lock(page, password, title):
    """The page encrypted with the password (AES-256-GCM, key from PBKDF2-SHA256), inside a small page
    that asks for it and decrypts it in the browser. Nothing readable is left in the file without the
    password; the encryption is standard, so it is as strong as the password is."""
    import base64
    import os

    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC

    salt, iv = os.urandom(16), os.urandom(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=PBKDF2_ROUNDS).derive(password.encode("utf-8"))
    sealed = AESGCM(key).encrypt(iv, page.encode("utf-8"), None)
    text = lambda b: base64.b64encode(b).decode("ascii")
    return (LOCKED_PAGE.replace("__TITLE__", html.escape(title)).replace("__SALT__", text(salt)).replace("__IV__", text(iv))
            .replace("__ROUNDS__", str(PBKDF2_ROUNDS)).replace("__DATA__", text(sealed)))


def ask_password():
    """The gallery password: from GALLERY_PASSWORD if set (for unattended runs), else typed twice."""
    import getpass
    import os

    if os.environ.get("GALLERY_PASSWORD"):
        return os.environ["GALLERY_PASSWORD"]
    while True:
        first = getpass.getpass("Password for the shared gallery (typing is hidden): ")
        if len(first) < 8:
            print("   Use at least 8 characters.")
            continue
        if getpass.getpass("Type it again: ") == first:
            return first
        print("   The two did not match. Try again.")


def to_clipboard(text):
    """Put a link on the clipboard: long links break when copied out of a wrapped console line."""
    import subprocess

    try:
        subprocess.run(["clip"], input=text.encode("utf-16-le"), check=True)
        return True
    except (OSError, subprocess.CalledProcessError):
        return False


def main():
    core.load_settings()
    base = Path(core.DEFAULT_SHOTS)
    if not base.is_dir():
        sys.exit(f"The screenshot folder does not exist yet: {base}\nRun Log_Screenshots.bat first.")
    print(f"Screenshot folder: {base}")
    print("Reading images and making thumbnails (only new or changed images take time)...")
    current = CURRENT_YEAR
    samples = collect(base, current)
    for year, paths in older_years():
        if year == current:
            continue
        shots = Path(paths["screenshot_folder"])
        print(f"Adding {year}: {shots}")
        older = collect(shots, year, paths["image_list"] or None, paths["synapse_folder"] or None,
                        root=core.share_url_for(shots).rstrip("/") + "/")
        print(f"   {len(older)} samples from {year}")
        samples += older
    if not samples:
        sys.exit("No images found. Run Log_Screenshots.bat first.")
    data = public_json(samples)
    page = PAGE.replace("__DATA__", data).replace("__BUILT__", html.escape(__import__("provenance").stamp(Path(__file__))))
    page = page.replace("__TITLE__", html.escape(GALLERY_TITLE))
    out = base / GALLERY_NAME
    out.write_text(page, encoding="utf-8")
    for old in OLD_NAMES:   # the page under a previous name would otherwise linger, out of date
        if old != GALLERY_NAME and (base / old).is_file():
            (base / old).unlink()
    print(f"\n{len(samples)} samples, {sum(len(s['pictures']) for s in samples)} images")
    print(f"Gallery: {out}")
    password = ask_password() if "--password" in sys.argv else None
    shared = None
    if "--server" in sys.argv or password:
        if Path(core.DEFAULT_OUT).is_dir():
            print("\nPacking the server copy (every image goes inside one file)" + (", locked with the password..." if password else "..."))
            copy, shared = server_copy(base, samples, password)
            print(f"Server copy: {copy}  ({copy.stat().st_size / 1e6:.0f} MB){'  PASSWORD PROTECTED' if password else ''}\nShare: {shared}")
        else:
            print(f"\nServer copy not made: cannot reach {core.DEFAULT_OUT}")
    if "--mirror" in sys.argv and Path(core.DEFAULT_OUT).is_dir():
        if password:
            # the mirror is the plain folders with full-size pictures: it would put back in the open what the
            # password is meant to keep out, so a protected run leaves it alone (an earlier mirror stays as it was)
            print("\nMirror NOT updated: it is not password protected and would publish the full-size pictures.")
        else:
            print("\nMirroring the gallery's folders to the server (only what changed is sent)...")
            shared = mirror(base)
            print(f"Full gallery on the server: {shared}")
    if shared and to_clipboard(shared):
        print("(The link to share is on the clipboard: paste it with Ctrl+V.)")
    return out


if __name__ == "__main__":
    path = main()
    if "--open" in sys.argv:
        import os
        os.startfile(path)
