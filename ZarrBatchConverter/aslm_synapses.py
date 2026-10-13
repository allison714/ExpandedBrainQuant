"""
Synapses and expansion factor in a stitched pan-ASLM volume (Neuroglancer precomputed, pan stain only).

  aslm_synapses.py <dataset folder> <tiles across> <tiles down> [overlap %] [number of depths]
  e.g.  aslm_synapses.py "X:\\...\\Stitched Precomputed NG\\260618_pan_aslm_4x4x9520" 4 4 15 6

The measurement is synapse_exf.py's, unchanged: the same finder, the same peak-to-peak reading
on a 2-pixel line, the same classes, the same log. What differs is where it looks:

  - The volume is far too large to go through whole, so it is sampled. A field of view is the
    middle FIELD_PX x FIELD_PX pixels of one tile on one plane: inside the tile, away from the
    bands where two tiles were blended. Every tile is sampled at each of the chosen depths.
  - The expansion factor is the mean and s.d. across those fields, as for the other samples.
  - There is no antibody, so synapses cannot be confirmed by one: every number here rests on
    double peaks in the pan stain alone, which read high (see the Method sheet).
  - Tiles differ in brightness, so the log is ordered by how bright each synapse is within its
    own field, not by raw counts; otherwise one bright tile would fill the log.

Finished fields are kept in <output>/_fields.pkl, so a run that is stopped carries on where it was.
"""

import csv
import json
import math
import pickle
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import numpy as np

import ab_efficiency as eff
import convert_to_zarr as core
import stitch_check
import synapse_exf as syn
import synapse_structure as structure

FIELD_PX = 2048          # side of one field of view, pixels
READERS = 16             # chunk files read at once
LOG_NAME = "pan-ASLM Synapse Log.xlsx"


def read_plane(vol, x0, x1, y0, y1, z):
    """One XY window of the full-resolution volume, its chunk files read side by side."""
    s = vol.scale(0)
    cx, cy, cz = s["chunk_sizes"][0]
    out = np.zeros((y1 - y0, x1 - x0), vol.dtype)
    z0 = (z // cz) * cz
    corners = [(gx, gy) for gx in range((x0 // cx) * cx, x1, cx) for gy in range((y0 // cy) * cy, y1, cy)]

    def one(corner):
        gx, gy = corner
        return gx, gy, vol.chunk(s, gx, gy, z0)[:, :, z - z0]

    with ThreadPoolExecutor(READERS) as pool:
        for gx, gy, block in pool.map(one, corners):
            ax, ay = max(gx, x0), max(gy, y0)
            bx, by = min(gx + block.shape[0], x1), min(gy + block.shape[1], y1)
            out[ay - y0:by - y0, ax - x0:bx - x0] = block[ax - gx:bx - gx, ay - gy:by - gy].T
    return out


def run(folder, nx, ny, overlap, depths):
    from scipy import ndimage

    core.load_settings()
    vol = stitch_check.Volume(folder)
    s0 = vol.scale(0)
    sx, sy, sz = s0["size"]
    um = s0["resolution"][0] / 1000
    name = next((p.name for p in (Path(folder), *Path(folder).parents) if len(p.name) > 6 and p.name[:6].isdigit()), Path(folder).name)
    base = Path(core.DEFAULT_SHOTS).parent / "Zarr Synapse ExF" / "pan-ASLM"
    out = base / name
    out.mkdir(parents=True, exist_ok=True)
    log = base / LOG_NAME
    try:
        kept = syn.read_decisions(log)
        if log.is_file():
            open(log, "ab").close()
    except PermissionError:
        sys.exit(f"{LOG_NAME} is open in Excel. Close it and run this again.")
    tile_x, tile_y = sx / (nx - (nx - 1) * overlap), sy / (ny - (ny - 1) * overlap)
    pitch_x, pitch_y = tile_x * (1 - overlap), tile_y * (1 - overlap)
    planes = sorted({int(sz * f) for f in np.linspace(0.08, 0.88, depths)})
    half = int(round(2 * syn.PROFILE_HALF_NM * syn.EXF_NOMINAL / (um * 1000) * 0.75))
    print(f"{name}: {sx} x {sy} x {sz} voxels, pixel {um * 1000:.0f} nm as imaged; {nx} x {ny} tiles, {len(planes)} depths: "
          f"{nx * ny * len(planes)} fields of {FIELD_PX} px\nOutput: {out}\n")

    done_file = out / "_fields.pkl"
    done = pickle.loads(done_file.read_bytes()) if done_file.is_file() else {}
    for z in planes:
        for j in range(ny):
            for i in range(nx):
                key = (z, j, i)
                if key in done:
                    continue
                t0 = time.time()
                x0 = int(i * pitch_x + tile_x / 2 - FIELD_PX / 2)
                y0 = int(j * pitch_y + tile_y / 2 - FIELD_PX / 2)
                x0, y0 = max(0, min(x0, sx - FIELD_PX)), max(0, min(y0, sy - FIELD_PX))
                field = read_plane(vol, x0, x0 + FIELD_PX, y0, y0 + FIELD_PX, z)
                hits = syn.find_synapses(field, None, um)
                smooth = ndimage.gaussian_filter(field.astype(np.float32), syn.SMOOTH_PX)
                for h in hits:
                    h["crop"] = syn.crop(smooth, h["row"], h["col"], half).copy()        # for its picture, so nothing is read twice
                    h["row"], h["col"], h["z"], h["fov"] = h["row"] + y0, h["col"] + x0, z, key
                    h.pop("lines", None)
                valid = field > 0
                done[key] = {"hits": hits, "area": int(valid.sum()), "mean": float(field[valid].mean()) if valid.any() else 0.0,
                             "p999": float(np.percentile(field[valid], 99.9)) if valid.any() else 0.0}
                done_file.write_bytes(pickle.dumps(done))
                print(f"   z {z:5d}  tile {i},{j}: {len(hits):3d} synapse profile(s), mean {done[key]['mean']:.1f} counts  ({time.time() - t0:.0f} s)", flush=True)

    fields = {k: v for k, v in done.items() if k[0] in planes}
    hits = syn.unique([h for v in fields.values() for h in v["hits"]], syn.REF_NM * syn.EXF_NOMINAL / (um * 1000))
    area = sum(v["area"] for v in fields.values())
    for h in hits:
        h["ab_pos"] = h["ab2_pos"] = False
        syn.classify(h, None, um)
        h["example"] = syn.example_score(h, um)
        structure.measure(h, um)
    hits.sort(key=lambda h: (-h["bright"], -h["r2"] * h["dip"]))      # brightest for its own field first, then the cleanest profile
    print(f"\n{len(hits)} synapse profile(s) in {len(fields)} fields")

    with open(out / f"{name}_fields.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Slice", "Depth (µm, as imaged)", "Tile Column", "Tile Row", "Synapses", "ExF (mean)", "ExF (s.d.)", "Mean Counts", "99.9th Percentile Counts"])
        for (z, j, i), v in sorted(fields.items()):
            e = [h["exf"] for h in hits if h["fov"] == (z, j, i)]
            w.writerow([z, round(z * s0["resolution"][2] / 1000, 1), i, j, len(e), round(float(np.mean(e)), 2) if e else "",
                        round(float(np.std(e, ddof=1)), 2) if len(e) > 1 else "", round(v["mean"], 1), round(v["p999"], 1)])
    with open(out / f"{name}_synapses.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Slice", "X (px)", "Y (px)", "DP-PSD peak to peak (µm, as imaged)", "ExF (peak to peak)", "ExF (Gaussian fit)", "Field of View",
                    "Class", "PSD Named By", "PSD Length (µm, as imaged)", "PSD Thickness (µm, as imaged)", "Orientation (deg)", "Fit R²", "Dip",
                    "Peak Ratio", "Brightness Rank"])
        for h in hits:
            w.writerow([h["z"], round(h["col"], 1), round(h["row"], 1), round(h["distance_um"], 4), round(h["exf"], 2), round(h["exf_fit"], 2),
                        "z{}-r{}-c{}".format(*h["fov"]), h["class"], h["psd_by"], "" if h["psd_len_um"] is None else round(h["psd_len_um"], 3),
                        round(h["psd_fwhm_um"], 3), round(math.degrees(h["theta"]) % 180, 1), round(h["r2"], 3), round(h["dip"], 3),
                        round(h["ratio"], 3), round(h["bright"], 4)])

    # ---- the log: pictures, rows and the summary, as for the other samples
    view = Path(folder).parent / core.STATE_DIR / f"{name}.json"
    state = json.loads(view.read_text(encoding="utf-8")) if view.is_file() else None
    pan_ch = {"label": "488 nm"}
    for stale in out.glob(f"{name}_synapse-*.png"):
        stale.unlink()
    logged, rows, shown, decisions = hits[:syn.TOP_N], [], [], []
    bar_exf = syn.sample_exf([h for h in logged if str(syn.decision_for(kept, name, h)[0] or "").strip().upper() == "Y"], [], hits, False).get("ExF")
    for number, h in enumerate(logged, start=1):
        tile = {"half": half, "pan": h["crop"], "ab": None}
        png = out / f"{name}_synapse-{number:03d}_z{h['z']:04d}_x{round(h['col']):05d}_y{round(h['row']):05d}.png"
        syn.card(png, number, name, h, tile, um, pan_ch, None, bar_exf)
        link = None
        if state:
            place = {"x": h["col"], "y": h["row"], "z": h["z"]}
            here = dict(state, position=[place[n] for n in state["dimensions"]], crossSectionScale=0.15, layout="xy")
            link = core.state_to_link(here, Path(folder).parent / core.STATE_DIR / "synapses" / name / f"synapse_{number:04d}.json")
        accept, notes = syn.decision_for(kept, name, h)
        decisions.append(accept)
        shown.append((number, h, tile))
        rows.append({
            "#": number, "Sample": name, "Slice": h["z"], "X (px)": round(h["col"]), "Y (px)": round(h["row"]), "Antibody": "none (pan only)",
            "DP–PSD (µm, as imaged)": round(h["distance_um"], 3), "ExF (peak to peak)": round(h["exf"], 1), "ExF (Gaussian fit)": round(h["exf_fit"], 1),
            "In 16–24": "Y" if syn.EXF_EXPECTED[0] <= h["exf"] <= syn.EXF_EXPECTED[1] else "no", "Fit R²": round(h["r2"], 3), "Dip": round(h["dip"], 2),
            "Peak Ratio": round(h["ratio"], 2), "Class": h["class"], "PSD Named By": h["psd_by"], "Example Score": round(h["example"], 2),
            "Synapse ID": f"z{h['z']:04d}_x{round(h['col']):05d}_y{round(h['row']):05d}", "Accept": accept, "Notes": notes,
            "Picture": str(png), "Neuroglancer": link})
    if shown:
        syn.contact_sheet(out / f"{name}_synapses_contact-sheet.png", name, shown, um, None, bar_exf)
    accepted = [h for h, d in zip(logged, decisions) if str(d or "").strip().upper() == "Y"]
    no = sum(str(d or "").strip().upper() == "N" for d in decisions)
    quoted = syn.sample_exf(accepted, [], hits, False)
    counted, counted_on = ((accepted, "accepted synapses") if len(accepted) >= syn.MIN_ACCEPTED else (hits, "all double peaks (includes membrane pairs)"))
    every = [h["exf"] for h in hits]
    exf = quoted.get("ExF")
    built = structure.summarise(counted, hits, area, um, quoted, out / f"{name}_synapses_structure.png", name, None) if hits else {}
    summary = [{
        **quoted, **built,
        **{k: sum(h["class"] == c for h in counted) for k, c in (("Type I Macular", "Type I, macular"), ("Type I Perforated", "Type I, perforated"),
                                                                 ("Type II (putative)", "Type II (putative)"))},
        "Classes Counted On": counted_on, "Pixel (nm, as imaged)": round(um * 1000, 1), "Pixel (nm, corrected)": round(um * 1000 / exf, 1) if exf else None,
        "Contact Sheet": str(out / f"{name}_synapses_contact-sheet.png") if shown else None, "Sample": name, "Antibody": "none (pan only)", "Synaptic": "no",
        "Planes": len(planes), "Synapses found": len(hits), "AB+": 0, "ExF, all found": syn.summarise(every), "ExF, AB+ only": syn.summarise([]),
        "ExF, logged": syn.summarise([h["exf"] for h in logged]), "ExF, accepted": syn.summarise([h["exf"] for h in accepted]),
        "Accepted / Rejected / Undecided": f"{len(accepted)} / {no} / {len(logged) - len(accepted) - no}",
        "Found outside 16–24 (%)": round(100 * np.mean([not syn.EXF_EXPECTED[0] <= e <= syn.EXF_EXPECTED[1] for e in every]), 1) if every else None}]
    # other volumes already in this log stay in it
    if log.is_file():
        from openpyxl import load_workbook

        old = load_workbook(log)
        for sheet, key, into in (("Synapses", syn.HEADER_ROW, rows), ("Summary", 1, summary)):
            names = [c.value for c in old[sheet][key]]
            for cells in old[sheet].iter_rows(min_row=key + 1):
                row = {h: (c.hyperlink.target if c.hyperlink else c.value) for h, c in zip(names, cells)}
                if row.get("Sample") and row["Sample"] != name:
                    into.append(row)
    syn.write_log(log, rows, summary)
    s = summary[0]
    print(f"   ExF quoted : {s.get('ExF')} ± {s.get('ExF s.d.')}, 95% CI {s.get('ExF 95% CI')}, {s.get('ExF at Window Edge (%)')}% at the window edge "
          f"({s.get('ExF Basis')}; fields {s.get('ExF Fields')}, n = {s.get('ExF n')}); Gaussian fit {s.get('ExF (Gaussian fit)')}\n"
          f"   all found  : {s['ExF, all found']}; outside 16–24: {s['Found outside 16–24 (%)']}%\n"
          f"   classes    : macular {s['Type I Macular']}, perforated {s['Type I Perforated']}, type II (putative) {s['Type II (putative)']} ({counted_on})\n"
          f"   PSD        : length {s.get('PSD Length (nm, corrected)')} nm, thickness {s.get('PSD Thickness (nm, corrected, with blur)')} nm\n"
          f"   anisotropy : {s.get('ExF Anisotropy (%)')}% (p = {s.get('Anisotropy p')}); field CV {s.get('ExF Field CV (%)')}%\n\nDone. {log}")


if __name__ == "__main__":
    run(sys.argv[1], int(sys.argv[2]), int(sys.argv[3]), float(sys.argv[4]) / 100 if len(sys.argv) > 4 else 0.15,
        int(sys.argv[5]) if len(sys.argv) > 5 else 6)
