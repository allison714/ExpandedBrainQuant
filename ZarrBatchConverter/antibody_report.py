"""Write the antibody report (SZ vs AC, WT vs KO) from the comparison CSVs and the synapse log.

Run compare_ab.py SZ AC and compare_ab.py WT KO first; this only reads their results.
"""
import csv
import html
import urllib.parse
from pathlib import Path

from openpyxl import load_workbook

root = Path(r"C:\Users\admin\OneDrive - Yale University\09_Data & Imaging\Images\2026")
comp = root / "Zarr Comparisons"
esc = html.escape
link = lambda rel: urllib.parse.quote(rel)


def read(folder, name):
    return list(csv.DictReader(open(comp / folder / name, encoding="utf-8-sig")))


def table(heads, rows, cls=""):
    out = [f'<table class="{cls}"><thead><tr>' + "".join(f"<th>{esc(h)}</th>" for h in heads) + "</tr></thead><tbody>"]
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{c if str(c).startswith('<') else esc(str(c))}</td>" for c in r) + "</tr>")
    return "\n".join(out) + "</tbody></table>"


D = "Puncta / 1000 µm² (as imaged)"
S = "Punctum Signal (noise-widths)"
C = "Punctum Signal (counts above background)"
A = "Labelled Area (%)"
tag = lambda v: f'<span class="v {"down" if v.startswith("DECREASED") else "up" if v.startswith("increased") else "same"}">{esc(v.lower())}</span>'

# ---- SZ vs AC
sz = read("SZ vs AC", "Comparison.csv")
szp = read("SZ vs AC", "Per Sample.csv")
t_sz = table(["Antibody", "SZ, 24 Feb 2026", "AC, 29 Apr 2026", "No-primary controls (SZ)", "AC against SZ"],
             [[r["Antibody"], r[f"{D}: SZ"], r[f"{D}: AC"], r["Control Floor (puncta / 1000 µm²)"], tag(r[f"{D}: Verdict"])] for r in sz])
dark = [r for r in szp if r["Set"] == "AC" or "control" in r["Antibody"] or "surface" in r["Sample"]]
t_dark = table(["Sample", "Set", "Puncta / 1000 µm²", "Background (counts)", "Noise (counts)"],
               [[r["Sample"], r["Set"] + (" control" if "control" in r["Antibody"] else ""), r[D], r["Background (counts)"], r["Noise (counts)"]] for r in dark], "small")

# ---- WT vs KO
ko = read("WT vs KO", "Comparison.csv")
kop = read("WT vs KO", "Per Sample.csv")
t_ko = table(["Antibody", "Measure", "WT", "KO", "KO / WT (95% interval)", "p", "KO against WT"],
             [[r["Antibody"] if m == D else "", m, r[f"{m}: WT"], r[f"{m}: KO"], r[f"{m}: KO / WT"], r[f"{m}: p"], tag(r[f"{m}: Verdict"])]
              for r in ko for m in (D, S, C, A)])
rows = load_workbook(root / "Zarr Synapse ExF" / "Synapse Log.xlsx", read_only=True)["Summary"].iter_rows(values_only=True)
heads = next(rows)
log = {r[0]: dict(zip(heads, r)) for r in rows}
g = lambda name, key: "–" if log.get(name, {}).get(key) is None else log[name][key]
t_syn = table(["Set", "Antibody", "Sample", "Type", "Synapses found", "With antibody", "Synapses labelled (%)", "Puncta on synapses (× chance)", "ExF"],
              [[r["Set"], r["Antibody"], r["Sample"], r["Type"], g(r["Sample"], "Synapses found"), g(r["Sample"], "AB+"),
                g(r["Sample"], "AB Labelled Synapses (%)"), g(r["Sample"], "AB Enrichment (× chance)"),
                (f"ExF = {g(r['Sample'], 'ExF')} ± {g(r['Sample'], 'ExF s.d.')}" if g(r["Sample"], "ExF") != "–" else "–")
                + (" *" if str(g(r["Sample"], "ExF Basis")).startswith("all double") else "")] for r in kop], "small")

page = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Antibody Labelling Report</title>
<style>
:root {{ --bg:#ffffff; --text:#1b1f24; --dim:#5b6470; --line:#d9dee5; --panel:#f5f7fa; --down:#b3261e; --up:#1a6f3c; --same:#5b6470; --accent:#0b5cad; }}
@media (prefers-color-scheme: dark) {{ :root {{ --bg:#111418; --text:#e6e9ee; --dim:#9aa4b2; --line:#2c333d; --panel:#1b2027; --down:#ff8a80; --up:#7ddc9c; --same:#9aa4b2; --accent:#6db3ff; }} }}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:var(--bg); color:var(--text); font:15px/1.55 Arial, "Segoe UI", sans-serif; }}
main {{ max-width:1100px; margin:0 auto; padding:28px 18px 60px; }}
h1 {{ font-size:26px; margin:0 0 4px; }} h2 {{ font-size:20px; margin:38px 0 8px; padding-top:14px; border-top:1px solid var(--line); }}
h3 {{ font-size:16px; margin:22px 0 6px; }}
.sub {{ color:var(--dim); margin:0 0 18px; }}
.box {{ background:var(--panel); border:1px solid var(--line); border-radius:8px; padding:12px 16px; margin:14px 0; }}
.box li {{ margin:5px 0; }}
table {{ border-collapse:collapse; width:100%; margin:10px 0 6px; font-size:14px; }}
th, td {{ text-align:left; padding:6px 10px; border-bottom:1px solid var(--line); vertical-align:top; }}
th {{ font-weight:600; background:var(--panel); }}
table.small {{ font-size:12.5px; }} table.small td, table.small th {{ padding:4px 8px; }}
.wrap {{ overflow-x:auto; }}
.v {{ font-weight:700; }} .down {{ color:var(--down); }} .up {{ color:var(--up); }} .same {{ color:var(--same); font-weight:400; }}
figure {{ margin:16px 0; }} figure img {{ max-width:100%; border:1px solid var(--line); border-radius:6px; display:block; }}
figcaption {{ color:var(--dim); font-size:13px; margin-top:6px; }}
a {{ color:var(--accent); }}
.note {{ color:var(--dim); font-size:13px; }}
@media print {{ h2 {{ break-before:page; }} }}
</style></head><body><main>
<h1>Antibody Labelling Report</h1>
<p class="sub">Kuan Lab pan-ExM · written 7 October 2026 · two comparisons: SZ against AC (single-label screen), and wild type against knockout (31 August 2026 set)</p>

<div class="box"><strong>In short</strong>
<ul>
<li><strong>SZ against AC:</strong> these are the same gels, imaged by SZ before the pan stain and by AC nine weeks later, after it. In the AC images the antibody channel holds no signal at all, for all five antibodies: fewer puncta than even the no-primary controls, with identical background and noise in every sample. Either the antibody signal did not survive the pan-staining and the nine weeks, or the 637 channel was imaged with too little excitation. The images cannot tell these apart.</li>
<li><strong>Wild type against Shank2 knockout:</strong> Cx43 labelling is about a fifth of wild type in the knockout animal. GluA1 (rabbit) and GluA3 (mouse) are not reduced. Each antibody was imaged in one animal per genotype, so none of this can yet be put down to the genotype.</li>
<li>Nothing here has been reviewed by a person or checked against hand measurements.</li>
</ul></div>

<p>Every picture of every sample is in the <a href="{link('../Zarr Screenshots/2026 AC Gallery.html')}">2026 AC Gallery</a>.</p>

<h2>1. SZ against AC: did labelling hold over nine weeks?</h2>
<p>These are the same gels imaged twice: by SZ on 24 February 2026, before pan-staining, and by AC on 29 April 2026, after the pan stain was added. Only the antibody channel (637 nm) is used, because the SZ images have no pan stain. Each number is the mean ± s.d. over fields of view (one camera frame each); n is the number of fields.</p>
<h3>Antibody puncta per 1000 µm² (as imaged)</h3>
<div class="wrap">{t_sz}</div>
<p class="note">A punctum is 3 or more connected pixels at least 6 noise-widths above the field's median. Bassoon was imaged by SZ only and is not compared.</p>

<figure><img src="{link('SZ vs AC/SZ vs AC - Examples.png')}" alt="Example fields, SZ above AC">
<figcaption>The busiest field of each sample, antibody channel, background removed, the same brightness scale within each column. SZ above, AC below.</figcaption></figure>

<h3>What the AC channel looks like, and the two explanations</h3>
<ul>
<li><strong>AC is far below the no-primary controls.</strong> The SZ controls, which have no primary antibody, still show 2 to 8 puncta per 1000 µm² from stray secondary antibody. AC shows hundreds of times fewer. Fading alone would leave the brightest aggregates; this is everything gone.</li>
<li><strong>The AC channel looks like a dark frame.</strong> Its background (110.0 counts) and noise (1.53 counts) are the same in all five AC samples, and lower than in the SZ controls.</li>
<li><strong>The other AC channel is fine.</strong> The 488 channel of the same files shows normal tissue structure.</li>
<li><strong>The channel preset differs:</strong> SZ used <code>637_Cy5_CF40</code>, AC used <code>637_CF25</code>. Objective (40×), pinhole reading (40) and binning (1 × 1) match.</li>
</ul>
<div class="wrap">{t_dark}</div>
<p><strong>Explanation 1: the label was lost.</strong> Pan-staining means more handling of the gel (staining, washes, re-expansion), and nine weeks passed. If that removed or destroyed the secondary antibody's dye, every punctum would go, including the stray secondary in the controls. That fits.</p>
<p><strong>Explanation 2: the channel was not excited.</strong> A 637 channel imaged with very low laser power or exposure gives exactly this picture as well. The preset is different, and laser power and exposure are not stored in the image files.</p>
<p><strong>To settle it:</strong> image one AC gel again with SZ's <code>637_Cy5_CF40</code> preset and settings. If the puncta are back, it was the imaging. If not, repeat on a fresh gel: image the antibody, pan-stain, and image again the same day, which separates the pan-staining from the nine weeks.</p>
<figure><img src="{link('SZ vs AC/SZ vs AC - Fields.png')}" alt="All fields, SZ and AC">
<figcaption>Every field of view as a dot, for four measures. Blue = SZ, orange = AC, white bar = mean, dashed line = no-primary controls.</figcaption></figure>

<h2>2. Wild type against Shank2 knockout (31 August 2026)</h2>
<p>The knockout animals (KO1b, KO2, KO3, KO6) lack Shank2. None of the antibodies here is against Shank2, so this section does not test whether an antibody is specific. It asks whether the labelling of Cx43, GluA1 and GluA3 differs in tissue without Shank2.</p>
<p>All samples were imaged on the same day with the same channel preset, so imaging differences are much less of a worry here. Single tiles (<code>_F0</code>) are left out. GluA2, GluA3 (rabbit) and mGluR5 exist in one genotype only and are not compared.</p>
<h3>From the antibody channel alone</h3>
<div class="wrap">{t_ko}</div>
<p class="note">"Decreased" or "increased" is said only when the whole 95% interval of the ratio is on one side of 1. p is from a Mann–Whitney test on the fields.</p>
<figure><img src="{link('WT vs KO/WT vs KO - Examples.png')}" alt="Example fields, WT above KO">
<figcaption>The busiest field of each, same brightness scale within each column. Wild type above, knockout below.</figcaption></figure>

<h3>At the synapses</h3>
<p>These samples have a pan stain, so synapses were found from it and each was checked for antibody. "Synapses labelled" is the share of synapses found that carry antibody. "Puncta on synapses" is how many times more often antibody puncta sit on a synapse than chance would put them there; 1 means no relation to synapses.</p>
<div class="wrap">{t_syn}</div>
<p class="note">* ExF from pan-only double peaks: reads high and should not be quoted. ExF is the DP–PSD distance read peak to peak ÷ 81.9 nm, mean ± s.d. across fields of view (M'Saad, Cairns et al., Nat. Biotechnol. 2025).</p>

<h3>Reading it</h3>
<ul>
<li><strong>Cx43</strong> is 0.19 of wild type in puncta density and 0.05 in labelled area in the Shank2 knockout animal. In wild type its puncta are not at synapses (1.2 to 1.7 times chance), as expected for an astrocytic gap-junction protein. One knockout sample is named <code>637at5per</code>, which suggests a different laser setting; leaving it in or out does not change the direction of the result.</li>
<li><strong>GluA1 (rabbit)</strong> is not reduced: 51 to 65% of synapses are labelled in the knockout against 45 to 65% in wild type, and puncta density is 1.38 times higher.</li>
<li><strong>GluA3 (mouse)</strong> is not reduced either, and labels few synapses in both genotypes (2 to 24%).</li>
</ul>
<div class="box"><strong>Why this is not yet a genotype effect.</strong>
<ul>
<li><strong>One animal per genotype.</strong> Cx43 was imaged in WT1 and KO1b, GluA1 in WT2 and KO2, GluA3 in WT6 and KO6. The fields of view are repeats within one animal and one gel, not independent animals. A difference between two gels (staining batch, depth imaged, region) would look exactly like this.</li>
<li><strong>The Cx43 result is the surprising one.</strong> A fivefold loss of an astrocytic protein in a Shank2 knockout would need more animals, and a check that the same brain region and depth were imaged, before it is believed.</li>
<li>The intervals and p-values describe how well each animal was sampled. They say nothing about animal-to-animal variation.</li>
</ul></div>
<figure><img src="{link('WT vs KO/WT vs KO - Fields.png')}" alt="All fields, WT and KO">
<figcaption>Every field of view as a dot. Blue = wild type, orange = knockout, white bar = mean.</figcaption></figure>

<h2>3. How it was measured</h2>
<ul>
<li><strong>Fields of view.</strong> Eight planes through the middle 80% of each stack, or the single plane of an overview; each camera-frame-sized tile of a plane is one field.</li>
<li><strong>Puncta.</strong> The antibody channel is smoothed (1-pixel Gaussian). A punctum is 3 or more connected pixels at least 6 noise-widths (1.4826 × MAD) above the field's median.</li>
<li><strong>Comparison.</strong> Ratio of the two sets' field means, with a bootstrap 95% interval (2,000 resamples of the fields within each set), and a Mann–Whitney test.</li>
<li><strong>Synapses.</strong> Found in the pan stain as a PSD with a parallel row of dense projections; antibody-positive when the antibody on the synapse is at least 6 noise-widths above the plane's median. The full method is on the Method sheet of the Synapse Log.</li>
</ul>

<h2>4. What to keep in mind</h2>
<ul>
<li>Laser power and exposure are not recorded in the files. Any comparison across imaging sessions (SZ against AC) rests on them having matched.</li>
<li>Fields from the same sample are not independent, so the p-values are optimistic. The size of the effect and its interval matter more.</li>
<li>SZ surface-to-surface stacks and single-plane overviews are pooled per antibody; the per-sample tables keep them apart.</li>
<li>Synapse finding sees only synapses lying side-on, and without a synaptic antibody some "synapses" are pairs of membranes. "Synapses labelled" is therefore a lower bound.</li>
<li>All of this is automatic and unreviewed.</li>
</ul>

<h2>5. Files</h2>
<ul>
<li><a href="{link('SZ vs AC/Comparison.csv')}">SZ vs AC/Comparison.csv</a>, <a href="{link('SZ vs AC/Per Sample.csv')}">Per Sample.csv</a>, <a href="{link('SZ vs AC/Imaging Settings.txt')}">Imaging Settings.txt</a></li>
<li><a href="{link('WT vs KO/Comparison.csv')}">WT vs KO/Comparison.csv</a>, <a href="{link('WT vs KO/Per Sample.csv')}">Per Sample.csv</a>, <a href="{link('WT vs KO/Imaging Settings.txt')}">Imaging Settings.txt</a></li>
<li><a href="{link('../Zarr Synapse ExF/Synapse Log.xlsx')}">Synapse Log.xlsx</a> (synapses, expansion factors, method)</li>
<li><a href="{link('../Zarr Screenshots/2026 AC Gallery.html')}">2026 AC Gallery</a></li>
</ul>
<p class="note">Made with compare_ab.py and synapse_exf.py in Desktop\\ZarrBatchConverter. To repeat: Compare_AB.bat SZ AC, and Compare_AB.bat WT KO.</p>
</main></body></html>
"""
out = comp / "Antibody Report.html"
out.write_text(page, encoding="utf-8")
print(out, len(page))
